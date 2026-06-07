from __future__ import annotations

import argparse
import io
import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "刘可毕业答辩_答辩优化.pptx"
DEFAULT_OUTPUT = ROOT / "刘可毕业答辩_答辩优化版.pptx"
DEFAULT_ASSET_DIR = ROOT / "outputs" / "defense_ppt" / "one_image_assets"

SLIDE_W = 2000
SLIDE_H = 960
BODY_BOX = (Inches(0.72), Inches(1.18), Inches(11.58), Inches(5.95))

BG = "#F6F8FB"
WHITE = "#FFFFFF"
NAVY = "#19324C"
BLUE = "#2F6EA9"
TEAL = "#2C8C99"
GREEN = "#2E8B57"
ORANGE = "#F28C28"
RED = "#C94F4F"
GRAY = "#728197"
LIGHT_BLUE = "#EAF3FB"
LIGHT_TEAL = "#E8F6F7"
LIGHT_GREEN = "#ECF7F0"
LIGHT_ORANGE = "#FFF4E8"
LIGHT_RED = "#FBECEC"
LINE = "#D7DEE8"
TEXT = "#203246"
MUTED = "#556476"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Convert defense slides to a one-image-per-slide layout.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="Input PPTX path.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Output PPTX path.")
    parser.add_argument("--asset-dir", type=Path, default=DEFAULT_ASSET_DIR, help="Generated body image directory.")
    return parser.parse_args()


def font_path(bold: bool = False) -> str:
    candidates = [
        "/mnt/c/Windows/Fonts/msyhbd.ttc" if bold else "/mnt/c/Windows/Fonts/msyh.ttc",
        "/mnt/c/Windows/Fonts/simhei.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return candidate
    raise FileNotFoundError("No usable font file found for PIL rendering.")


FONT_REG = font_path(False)
FONT_BOLD = font_path(True)


def get_font(size: int, *, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size)


def remove_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def new_canvas() -> Image.Image:
    img = Image.new("RGB", (SLIDE_W, SLIDE_H), BG)
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle((18, 18, SLIDE_W - 18, SLIDE_H - 18), radius=40, outline="#E7ECF3", width=2)
    draw.ellipse((-160, -120, 260, 260), fill="#EDF4FB")
    draw.ellipse((SLIDE_W - 250, -120, SLIDE_W + 120, 220), fill="#F1F7FD")
    return img


def rounded_box(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], *, fill: str = WHITE, outline: str = LINE, width: int = 2, radius: int = 28) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def measure(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont) -> tuple[int, int]:
    if not text:
        return (0, 0)
    box = draw.multiline_textbbox((0, 0), text, font=font, spacing=6)
    return (box[2] - box[0], box[3] - box[1])


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, max_width: int) -> str:
    if not text:
        return ""
    tokens = re.findall(r"[A-Za-z0-9_./:+-]+|[^\x00-\x7F]| +", text)
    lines: list[str] = []
    current = ""
    for token in tokens:
        candidate = token if not current else current + token
        width, _ = measure(draw, candidate, font)
        if width <= max_width or not current:
            current = candidate
            continue
        lines.append(current.rstrip())
        current = token.lstrip()
    if current.strip():
        lines.append(current.rstrip())
    return "\n".join(lines)


def draw_wrapped_text(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    text: str,
    *,
    font: ImageFont.FreeTypeFont,
    fill: str = TEXT,
    max_width: int,
    spacing: int = 8,
) -> tuple[int, int]:
    wrapped = wrap_text(draw, text, font, max_width)
    draw.multiline_text(xy, wrapped, font=font, fill=fill, spacing=spacing)
    return measure(draw, wrapped, font)


def draw_card(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    *,
    title: str,
    body: str = "",
    accent: str = BLUE,
    fill: str = WHITE,
    title_size: int = 34,
    body_size: int = 25,
) -> None:
    rounded_box(draw, box, fill=fill, outline=LINE)
    x1, y1, x2, y2 = box
    draw.rounded_rectangle((x1 + 18, y1 + 18, x1 + 130, y1 + 34), radius=10, fill=accent)
    title_font = get_font(title_size, bold=True)
    body_font = get_font(body_size)
    draw_wrapped_text(draw, (x1 + 28, y1 + 52), title, font=title_font, fill=NAVY, max_width=x2 - x1 - 56, spacing=6)
    if body:
        draw_wrapped_text(draw, (x1 + 28, y1 + 112), body, font=body_font, fill=MUTED, max_width=x2 - x1 - 56, spacing=10)


def draw_pill(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    *,
    text: str,
    fill: str = LIGHT_BLUE,
    ink: str = NAVY,
    border: str = LINE,
    size: int = 24,
) -> None:
    rounded_box(draw, box, fill=fill, outline=border, width=1, radius=22)
    font = get_font(size, bold=True)
    w, h = measure(draw, text, font)
    x1, y1, x2, y2 = box
    draw.text((x1 + (x2 - x1 - w) / 2, y1 + (y2 - y1 - h) / 2 - 2), text, font=font, fill=ink)


def draw_bullet_list(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    items: list[str],
    *,
    font_size: int = 26,
    bullet_color: str = BLUE,
    text_color: str = MUTED,
    line_gap: int = 16,
) -> None:
    font = get_font(font_size)
    x1, y1, x2, _ = box
    y = y1
    for item in items:
        draw.ellipse((x1, y + 10, x1 + 12, y + 22), fill=bullet_color)
        _, h = draw_wrapped_text(draw, (x1 + 28, y), item, font=font, fill=text_color, max_width=x2 - x1 - 28, spacing=8)
        y += h + line_gap


def fit_image(src: Image.Image, size: tuple[int, int]) -> Image.Image:
    img = src.convert("RGB")
    return ImageOps.contain(img, size)


def paste_image(canvas: Image.Image, src: Image.Image, box: tuple[int, int, int, int], *, fill: str = WHITE, outline: str = LINE, caption: str | None = None) -> None:
    draw = ImageDraw.Draw(canvas)
    rounded_box(draw, box, fill=fill, outline=outline)
    x1, y1, x2, y2 = box
    caption_h = 56 if caption else 0
    inner = (x1 + 22, y1 + 22, x2 - 22, y2 - 22 - caption_h)
    fitted = fit_image(src, (inner[2] - inner[0], inner[3] - inner[1]))
    px = inner[0] + (inner[2] - inner[0] - fitted.width) // 2
    py = inner[1] + (inner[3] - inner[1] - fitted.height) // 2
    canvas.paste(fitted, (px, py))
    if caption:
        cap_font = get_font(24, bold=True)
        draw_wrapped_text(draw, (x1 + 28, y2 - 46), caption, font=cap_font, fill=MUTED, max_width=x2 - x1 - 56, spacing=6)


def extract_picture_shapes(slide) -> list[Image.Image]:
    images: list[Image.Image] = []
    for shape in slide.shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        blob = shape.image.blob
        images.append(Image.open(io.BytesIO(blob)).convert("RGB"))
    return images


def render_slide_02() -> Image.Image:
    img = new_canvas()
    draw = ImageDraw.Draw(img)
    banner = (70, 60, 1930, 180)
    rounded_box(draw, banner, fill=WHITE)
    draw_card(draw, banner, title="汇报结构", body="整体按“问题提出 - 方法设计 - 实验验证 - 系统演示 - 结论展望”展开。", accent=BLUE, fill=WHITE, title_size=40, body_size=28)
    sections = [
        ("01", "研究背景与问题定义", "图像传播场景、版权保护需求与本文目标。", BLUE, LIGHT_BLUE),
        ("02", "方法设计与模型结构", "Y 通道载荷、残差嵌入、恢复链路与损失函数。", TEAL, LIGHT_TEAL),
        ("03", "实验结果与对比分析", "主实验、Plain vs Robust 对比实验和消融实验。", GREEN, LIGHT_GREEN),
        ("04", "系统演示", "基于 Streamlit 的交互式水印嵌入与恢复演示。", ORANGE, LIGHT_ORANGE),
        ("05", "结论与展望", "本文结论、当前不足和后续改进方向。", RED, LIGHT_RED),
    ]
    card_w = 348
    gap = 22
    x = 70
    for num, title, body, accent, fill in sections:
        draw_card(draw, (x, 240, x + card_w, 820), title=num, body=f"{title}\n{body}", accent=accent, fill=fill, title_size=46, body_size=28)
        x += card_w + gap
    return img


def render_slide_03() -> Image.Image:
    img = new_canvas()
    draw = ImageDraw.Draw(img)
    cards = [
        ((70, 78, 965, 360), "应用场景", "社交媒体、在线图库和 AIGC 内容传播链路中，图像会被频繁复制、压缩、裁剪和二次编辑。", BLUE, WHITE),
        ((1035, 78, 1930, 360), "核心矛盾", "不可见性、可恢复性和鲁棒性难以同时兼顾，传统规则式方法在复杂攻击下容易退化。", ORANGE, WHITE),
        ((70, 400, 965, 682), "论文目标", "在保证宿主图像视觉质量的前提下，让水印可以稳定恢复，并对常见攻击保持一定鲁棒性。", TEAL, WHITE),
        ((1035, 400, 1930, 682), "研究价值", "既强调模型设计，也补齐训练、推理、批测、对比实验和消融实验的工程链路。", GREEN, WHITE),
    ]
    for box, title, body, accent, fill in cards:
        draw_card(draw, box, title=title, body=body, accent=accent, fill=fill)
    draw_card(draw, (70, 730, 1930, 890), title="一句话概括", body="数字图像版权保护需要“平时看不见、需要时能取出”的技术手段，而真实传播链路往往还会叠加缩放、模糊和马赛克等多类失真。", accent=BLUE, fill=LIGHT_BLUE, title_size=34, body_size=28)
    return img


def render_slide_04(pictures: list[Image.Image]) -> Image.Image:
    img = new_canvas()
    draw = ImageDraw.Draw(img)
    left_boxes = [
        ((70, 90, 760, 285), "问题一", "在保证视觉质量的前提下，把版权信息隐藏进宿主图像。", BLUE),
        ((70, 315, 760, 510), "问题二", "面对裁剪、JPEG、缩放、模糊和马赛克等攻击，恢复主要水印结构。", TEAL),
        ((70, 540, 760, 735), "问题三", "把训练、推理、批量测试、对比实验和消融实验串成可复现流程。", ORANGE),
    ]
    for box, title, body, accent in left_boxes:
        draw_card(draw, box, title=title, body=body, accent=accent)
    draw_card(draw, (800, 90, 1930, 270), title="核心思路", body="只编码水印模板的 Y 通道亮度结构，将其压缩成单通道噪声载荷，再通过 RGB 残差嵌入网络写入宿主图像。", accent=RED, fill=LIGHT_RED, title_size=34, body_size=28)
    if pictures:
        paste_image(img, pictures[0], (800, 320, 1930, 890), caption="基础框架示意图：在通用编码器 - 攻击层 - 解码器框架上，进一步落到 Y 通道载荷与 RGB 残差嵌入。")
    return img


def render_slide_05(pictures: list[Image.Image]) -> Image.Image:
    img = new_canvas()
    draw = ImageDraw.Draw(img)
    if pictures:
        paste_image(img, pictures[0], (70, 72, 1320, 890), caption="总体方案流程图")
    boxes = [
        ((1360, 120, 1930, 330), "输入侧", "宿主图像走 RGB，水印模板先转 YCbCr，并只保留 Y 通道参与学习。", BLUE),
        ((1360, 380, 1930, 590), "嵌入侧", "编码后的 payload 与 cover 一起进入 EmbedUNet，预测低幅值 RGB residual。", TEAL),
        ((1360, 640, 1930, 850), "恢复侧", "由 container 先恢复 payload，再解码回 watermark Y，完成版权信息读取。", GREEN),
    ]
    for box, title, body, accent in boxes:
        draw_card(draw, box, title=title, body=body, accent=accent)
    return img


def render_slide_06() -> Image.Image:
    img = new_canvas()
    draw = ImageDraw.Draw(img)
    draw_card(draw, (70, 60, 1930, 180), title="方法链路概括", body="把彩色水印转成亮度结构，再写成受控残差，最后通过两阶段恢复读回 watermark Y。", accent=BLUE, fill=LIGHT_BLUE, title_size=36, body_size=28)
    steps = [
        ("输入表示", "先抽取水印亮度结构，降低需要写入和恢复的信息复杂度。", BLUE, LIGHT_BLUE),
        ("载荷编码", "把 watermark Y 映射到单通道噪声载荷 payload。", TEAL, LIGHT_TEAL),
        ("残差嵌入", "通过受限残差写入版权信息，核心目标是减小宿主图像可感知失真。", ORANGE, LIGHT_ORANGE),
        ("含水印图生成", "残差与宿主图像相加后裁剪到合法像素范围。", GREEN, LIGHT_GREEN),
        ("恢复解码", "先从 container 中恢复 payload，再解码为 watermark Y。", RED, LIGHT_RED),
    ]
    positions = [
        (70, 240, 650, 470),
        (710, 240, 1290, 470),
        (1350, 240, 1930, 470),
        (390, 530, 970, 760),
        (1030, 530, 1610, 760),
    ]
    for (title, body, accent, fill), box in zip(steps, positions):
        draw_card(draw, box, title=title, body=body, accent=accent, fill=fill)
    chain_box = (290, 810, 1710, 900)
    rounded_box(draw, chain_box, fill=WHITE)
    draw_pill(draw, (340, 832, 610, 882), text="Watermark Y", fill=LIGHT_BLUE)
    draw_pill(draw, (645, 832, 915, 882), text="Payload", fill=LIGHT_TEAL)
    draw_pill(draw, (950, 832, 1280, 882), text="Residual Embed", fill=LIGHT_ORANGE)
    draw_pill(draw, (1315, 832, 1545, 882), text="Container", fill=LIGHT_GREEN)
    draw_pill(draw, (1580, 832, 1860, 882), text="Recover Y", fill=LIGHT_RED)
    arrow_font = get_font(30, bold=True)
    for x in (617, 922, 1287, 1552):
        draw.text((x, 838), "→", font=arrow_font, fill=GRAY)
    return img


def render_side_image_slide(pictures: list[Image.Image], *, cards: list[tuple[str, str, str]], caption: str) -> Image.Image:
    img = new_canvas()
    draw = ImageDraw.Draw(img)
    if pictures:
        paste_image(img, pictures[0], (70, 92, 1180, 890), caption=caption)
    y = 116
    fills = {
        BLUE: LIGHT_BLUE,
        TEAL: LIGHT_TEAL,
        GREEN: LIGHT_GREEN,
        ORANGE: LIGHT_ORANGE,
        RED: LIGHT_RED,
    }
    for title, body, accent in cards:
        draw_card(draw, (1225, y, 1930, y + 220), title=title, body=body, accent=accent, fill=fills.get(accent, WHITE))
        y += 250
    return img


def render_slide_07(pictures: list[Image.Image]) -> Image.Image:
    cards = [
        ("编码", "将 watermark Y 映射为单通道 payload。", BLUE),
        ("解码", "由恢复出的 payload 重建水印亮度结构。", TEAL),
        ("设计重点", "只恢复 Y 通道，减少颜色信息带来的额外难度，让模型优先学习稳定的亮度结构。", GREEN),
    ]
    return render_side_image_slide(pictures, cards=cards, caption="水印 Y 通道与噪声载荷 payload 之间的编码/解码结构")


def render_slide_08(pictures: list[Image.Image]) -> Image.Image:
    cards = [
        ("输入", "宿主图像和 payload 在通道维度联合输入。", BLUE),
        ("输出", "输出受限 RGB residual，避免明显破坏宿主图像。", TEAL),
        ("设计重点", "嵌入阶段写入的是残差而不是直接重绘图像，因此更容易控制不可见性。", ORANGE),
    ]
    return render_side_image_slide(pictures, cards=cards, caption="EmbedUNet 结合宿主图像和 payload，预测低幅值 RGB residual")


def render_slide_09(pictures: list[Image.Image]) -> Image.Image:
    cards = [
        ("恢复", "从 container 中估计 payload。", BLUE),
        ("攻击后", "攻击路径下仍尝试恢复主要 payload 结构。", TEAL),
        ("设计重点", "恢复链路先保证 payload 稳定，再通过解码器恢复水印，便于分别约束中间表征和最终水印质量。", RED),
    ]
    return render_side_image_slide(pictures, cards=cards, caption="从含水印图像中恢复 payload，再交给解码器重建 watermark Y")


def render_slide_10(pictures: list[Image.Image]) -> Image.Image:
    img = new_canvas()
    draw = ImageDraw.Draw(img)
    if pictures:
        paste_image(img, pictures[0], (70, 82, 1930, 640), caption="攻击仿真与鲁棒训练路径结构图")
    pills = [
        ("Crop", LIGHT_BLUE),
        ("JPEG", LIGHT_TEAL),
        ("Resize", LIGHT_GREEN),
        ("Blur", LIGHT_ORANGE),
        ("Mosaic", LIGHT_RED),
    ]
    x = 140
    for text, fill in pills:
        draw_pill(draw, (x, 700, x + 240, 760), text=text, fill=fill)
        x += 320
    draw_card(draw, (70, 790, 1930, 900), title="讲解重点", body="鲁棒训练让模型在训练阶段就接触真实传播失真。代价是 clean 条件下可能略有下降，但攻击后的恢复会更稳定。", accent=BLUE, fill=WHITE, title_size=32, body_size=28)
    return img


def render_slide_11() -> Image.Image:
    img = new_canvas()
    draw = ImageDraw.Draw(img)
    top = [
        ((70, 90, 620, 430), "MSE", "衡量像素级误差。数值越小，说明两张图在逐像素层面越接近；它也是 PSNR 的基础。", BLUE, LIGHT_BLUE),
        ((725, 90, 1275, 430), "PSNR", "把误差转换成分贝尺度。数值越高，通常表示失真越小，更适合横向比较图像质量。", TEAL, LIGHT_TEAL),
        ((1380, 90, 1930, 430), "SSIM", "衡量亮度、对比度和结构一致性，比单独看像素误差更接近人眼感受。", ORANGE, LIGHT_ORANGE),
    ]
    bottom = [
        ((70, 500, 620, 860), "Host 指标", "比较 cover 与 container，用来说明水印嵌入后宿主图像是否仍然自然、不可见。", GREEN, LIGHT_GREEN),
        ((725, 500, 1275, 860), "WM 指标", "比较 watermark Y 与 recovered watermark Y，用来说明水印亮度结构是否被稳定取回。", BLUE, WHITE),
        ((1380, 500, 1930, 860), "Attack 指标", "在攻击后重新计算 WM PSNR/SSIM，用来体现方法对压缩、裁剪和模糊等扰动的鲁棒性。", RED, LIGHT_RED),
    ]
    for box, title, body, accent, fill in top + bottom:
        draw_card(draw, box, title=title, body=body, accent=accent, fill=fill)
    return img


def render_slide_12() -> Image.Image:
    img = new_canvas()
    draw = ImageDraw.Draw(img)
    draw_card(draw, (70, 60, 1930, 182), title="总目标", body="同时约束不可见性、结构感知、payload 可恢复性和噪声分布稳定性。", accent=BLUE, fill=LIGHT_BLUE, title_size=38, body_size=29)
    cards = [
        ("Host Loss", "直接用 MSE 约束 cover 与 container 接近，是最基础的不可见性项。", BLUE, WHITE),
        ("HVS Loss", "在平坦区域给予更高惩罚，让嵌入更符合人眼视觉敏感性。", TEAL, WHITE),
        ("Payload Loss", "保证恢复链路先能把中间 payload 稳定取回。", GREEN, WHITE),
        ("Watermark Loss", "同时约束像素误差和结构相似性，避免只得到模糊但数值接近的恢复结果。", ORANGE, LIGHT_ORANGE),
        ("Noise Regularization", "限制 payload 的均值、方差和局部相关性，使其更接近噪声化分布。", RED, LIGHT_RED),
    ]
    positions = [
        (70, 240, 620, 520),
        (725, 240, 1275, 520),
        (1380, 240, 1930, 520),
        (400, 590, 950, 870),
        (1050, 590, 1600, 870),
    ]
    for (title, body, accent, fill), box in zip(cards, positions):
        draw_card(draw, box, title=title, body=body, accent=accent, fill=fill, title_size=30, body_size=25)
    return img


def render_slide_13(pictures: list[Image.Image]) -> Image.Image:
    img = new_canvas()
    draw = ImageDraw.Draw(img)
    left_cards = [
        ((70, 86, 500, 280), "数据集", "训练集使用 DIV2K；测试集为 inputs/chaos_batch 中的自建图像样本。", BLUE),
        ((530, 86, 960, 280), "实现框架", "PyTorch + 自定义攻击模块，代码已包含训练、推理和批量测试流程。", TEAL),
        ((70, 310, 500, 504), "评价对象", "Host 指标用于衡量不可见性；WM 指标用于衡量水印恢复质量。", GREEN),
        ((530, 310, 960, 504), "实验类型", "主实验、Plain vs Robust 对比、A0~A5 消融，以及攻击后样例可视化。", ORANGE),
    ]
    for box, title, body, accent in left_cards:
        draw_card(draw, box, title=title, body=body, accent=accent, title_size=30, body_size=24)
    draw_card(draw, (70, 550, 960, 890), title="代码链路", body="", accent=RED, fill=LIGHT_RED, title_size=32, body_size=24)
    mono_font = get_font(22)
    bullets = [
        "experiments/edge_wm_noise_train.py：训练主脚本",
        "experiments/edge_wm_noise_infer.py：单图推理与可视化导出",
        "experiments/edge_wm_noise_test.py：批量测试与指标汇总",
        "scripts/run_compare_robust_vs_plain.sh：对比实验",
        "scripts/run_ablation_suite.sh：消融实验",
        "app_demo.py：交互演示程序",
    ]
    draw_bullet_list(draw, (110, 634, 915, 860), bullets, font_size=22, bullet_color=RED)
    if pictures:
        paste_image(img, pictures[0], (1010, 86, 1930, 890), caption="主实验模型训练到位性诊断图")
    return img


def render_slide_14(pictures: list[Image.Image]) -> Image.Image:
    img = new_canvas()
    draw = ImageDraw.Draw(img)
    if len(pictures) >= 1:
        paste_image(img, pictures[0], (70, 78, 950, 620), caption="图 6-1 水印嵌入与恢复效果对比")
    if len(pictures) >= 2:
        paste_image(img, pictures[1], (1050, 78, 1930, 620), caption="图 6-2 常见攻击后的恢复效果")
    draw_card(draw, (70, 700, 950, 880), title="观察 1", body="clean 条件下宿主图几乎看不到明显水印痕迹，说明 residual 嵌入较克制。", accent=BLUE, fill=LIGHT_BLUE, title_size=32, body_size=28)
    draw_card(draw, (1050, 700, 1930, 880), title="观察 2", body="攻击会降低恢复清晰度，但多数场景下仍保留主要轮廓，与定量指标趋势一致。", accent=ORANGE, fill=LIGHT_ORANGE, title_size=32, body_size=28)
    return img


def render_slide_15(pictures: list[Image.Image]) -> Image.Image:
    img = new_canvas()
    draw = ImageDraw.Draw(img)
    metrics = [
        ("Clean Host PSNR", "40.013 dB", BLUE, LIGHT_BLUE),
        ("Clean Host SSIM", "0.9633", TEAL, LIGHT_TEAL),
        ("Clean WM PSNR", "28.595 dB", GREEN, LIGHT_GREEN),
        ("Clean WM SSIM", "0.9852", ORANGE, LIGHT_ORANGE),
        ("Attack WM PSNR", "12.807 dB", RED, LIGHT_RED),
        ("Attack WM SSIM", "0.8272", BLUE, WHITE),
    ]
    positions = [
        (70, 62, 640, 198),
        (715, 62, 1285, 198),
        (1360, 62, 1930, 198),
        (70, 220, 640, 356),
        (715, 220, 1285, 356),
        (1360, 220, 1930, 356),
    ]
    for (label, value, accent, fill), box in zip(metrics, positions):
        rounded_box(draw, box, fill=fill, outline=LINE)
        draw.text((box[0] + 28, box[1] + 22), label, font=get_font(24, bold=True), fill=MUTED)
        draw.text((box[0] + 28, box[1] + 62), value, font=get_font(36, bold=True), fill=NAVY)
    if len(pictures) >= 1:
        paste_image(img, pictures[0], (70, 408, 950, 810), caption="主实验模型最终测试质量对比")
    if len(pictures) >= 2:
        paste_image(img, pictures[1], (1050, 408, 1930, 810), caption="clean 与 attack-all 水印指标对比")
    draw_card(draw, (70, 842, 1930, 900), title="结论", body="主实验说明该方法在不可见性与可恢复性之间取得了平衡；综合攻击会显著压低恢复质量，但模型仍保留一定鲁棒性。", accent=BLUE, fill=WHITE, title_size=30, body_size=26)
    return img


def render_slide_16(pictures: list[Image.Image]) -> Image.Image:
    img = new_canvas()
    draw = ImageDraw.Draw(img)
    draw_card(draw, (70, 60, 900, 176), title="对比主旨", body="普通训练更偏向 clean 指标，鲁棒训练则在攻击后的稳定恢复上略有优势。", accent=BLUE, fill=LIGHT_BLUE, title_size=34, body_size=28)
    draw_card(draw, (1040, 60, 1930, 176), title="答辩可强调", body="这里体现的是“保真度 - 鲁棒性”的性能折中，而不是所有指标同时提升。", accent=ORANGE, fill=LIGHT_ORANGE, title_size=34, body_size=28)
    if len(pictures) >= 1:
        paste_image(img, pictures[0], (70, 220, 1000, 760), caption="普通训练与鲁棒训练的 clean/attack 核心指标对比")
    if len(pictures) >= 2:
        paste_image(img, pictures[1], (1070, 220, 1930, 760), caption="攻击后水印质量保持率")
    draw_card(draw, (70, 810, 1930, 900), title="结论", body="Plain 在 clean 条件下更优；Robust 在 attack WM PSNR 上略高，说明鲁棒训练牺牲部分 clean 表现，以换取攻击场景下更稳定的恢复。", accent=RED, fill=WHITE, title_size=30, body_size=26)
    return img


def render_slide_17(pictures: list[Image.Image]) -> Image.Image:
    img = new_canvas()
    draw = ImageDraw.Draw(img)
    if len(pictures) >= 1:
        paste_image(img, pictures[0], (70, 78, 950, 620), caption="水印 PSNR 均值与 P10-P90 区间")
    if len(pictures) >= 2:
        paste_image(img, pictures[1], (1050, 78, 1930, 620), caption="水印 SSIM 均值与 P10-P90 区间")
    findings = [
        ((70, 700, 620, 880), "A0_Full", "整体最均衡，说明完整模块与损失组合是合理的。", BLUE, LIGHT_BLUE),
        ((725, 700, 1275, 880), "A1_NoRobust", "下降最明显，说明鲁棒训练不仅影响攻击场景，也影响整体恢复稳定性。", RED, LIGHT_RED),
        ((1380, 700, 1930, 880), "A2_NoNoiseReg", "Attack-all 下 SSIM 下滑更明显，说明噪声正则有助于结构一致性。", ORANGE, LIGHT_ORANGE),
    ]
    for box, title, body, accent, fill in findings:
        draw_card(draw, box, title=title, body=body, accent=accent, fill=fill, title_size=29, body_size=25)
    return img


def render_slide_18(pictures: list[Image.Image]) -> Image.Image:
    img = new_canvas()
    draw = ImageDraw.Draw(img)
    if len(pictures) >= 1:
        paste_image(img, pictures[0], (70, 78, 950, 620), caption="相对 A0_Full 的 attack 指标变化")
    if len(pictures) >= 2:
        paste_image(img, pictures[1], (1050, 78, 1930, 620), caption="各消融变体的宿主图像质量")
    findings = [
        ((70, 700, 620, 880), "A3_NoHVS", "clean 水印上恢复较高，但并没有带来更好的综合表现。", BLUE, LIGHT_BLUE),
        ((725, 700, 1275, 880), "A4_NoPayload", "Host SSIM 很高，但水印恢复明显失衡，说明 payload 约束不能去掉。", RED, LIGHT_RED),
        ((1380, 700, 1930, 880), "A5_NoSSIMinWM", "结构相似性下降，仍需 SSIM 监督。", ORANGE, LIGHT_ORANGE),
    ]
    for box, title, body, accent, fill in findings:
        draw_card(draw, box, title=title, body=body, accent=accent, fill=fill, title_size=28, body_size=24)
    return img


def render_slide_19(pictures: list[Image.Image]) -> Image.Image:
    img = new_canvas()
    draw = ImageDraw.Draw(img)
    draw_card(draw, (70, 70, 600, 890), title="演示流程", body="", accent=BLUE, fill=LIGHT_BLUE, title_size=36, body_size=28)
    steps = [
        "上传载体图片",
        "选择 checkpoint 和水印模板",
        "设置攻击模式并开始演示",
        "观察含水印图、攻击后图和恢复水印",
    ]
    draw_bullet_list(draw, (115, 178, 548, 520), steps, font_size=28, bullet_color=BLUE, line_gap=18)
    draw_card(draw, (100, 610, 570, 850), title="答辩建议", body="现场优先播放录制好的 Streamlit 演示；如果设备不稳定，右侧四张结果图可以直接作为备用说明。", accent=ORANGE, fill=WHITE, title_size=30, body_size=26)
    labels = ["原始图片", "含水印图", "攻击后图", "攻击后恢复水印"]
    boxes = [
        (660, 82, 1260, 468),
        (1310, 82, 1930, 468),
        (660, 520, 1260, 890),
        (1310, 520, 1930, 890),
    ]
    for box, label, picture in zip(boxes, labels, pictures[:4]):
        paste_image(img, picture, box, caption=label)
    return img


def render_slide_20() -> Image.Image:
    img = new_canvas()
    draw = ImageDraw.Draw(img)
    conclusions = [
        ("结论 1", "Y 通道噪声载荷 + RGB 残差嵌入是一条可行路线。", BLUE, LIGHT_BLUE),
        ("结论 2", "主实验表明模型在不可见性与恢复质量之间取得了较好平衡。", TEAL, LIGHT_TEAL),
        ("结论 3", "鲁棒训练和复合损失能提升复杂攻击下的稳定恢复能力。", GREEN, LIGHT_GREEN),
        ("结论 4", "工程链路完整，可支撑复现与系统展示。", ORANGE, LIGHT_ORANGE),
    ]
    positions = [
        (70, 72, 505, 250),
        (545, 72, 980, 250),
        (1020, 72, 1455, 250),
        (1495, 72, 1930, 250),
    ]
    for (title, body, accent, fill), box in zip(conclusions, positions):
        draw_card(draw, box, title=title, body=body, accent=accent, fill=fill, title_size=28, body_size=24)
    draw_card(draw, (70, 315, 930, 890), title="当前不足", body="", accent=RED, fill=LIGHT_RED, title_size=36, body_size=28)
    draw_bullet_list(
        draw,
        (115, 430, 885, 835),
        [
            "强裁剪、强压缩和多攻击叠加下，恢复质量仍会明显下降。",
            "当前主要围绕固定模板和固定分辨率进行实验，泛化能力仍需扩展。",
            "指标仍以 PSNR/SSIM 为主，后续可加入误码率和主观评价。",
        ],
        font_size=27,
        bullet_color=RED,
        line_gap=26,
    )
    draw_card(draw, (1000, 315, 1930, 890), title="后续工作", body="", accent=BLUE, fill=WHITE, title_size=36, body_size=28)
    draw_bullet_list(
        draw,
        (1045, 430, 1880, 835),
        [
            "更真实的攻击仿真：屏摄、旋转、透视变换、链式压缩。",
            "更强的鲁棒训练策略：课程学习、自适应攻击采样、对抗式增强。",
            "在网络结构方面，尝试引入 CBAM 等轻量注意力机制，强化对边缘、纹理和水印相关区域的特征建模。",
            "更广的工程目标：多模板、多载荷、可变分辨率和真实应用部署。",
        ],
        font_size=27,
        bullet_color=BLUE,
        line_gap=24,
    )
    return img


RENDERERS = {
    2: lambda pictures: render_slide_02(),
    3: lambda pictures: render_slide_03(),
    4: render_slide_04,
    5: render_slide_05,
    6: lambda pictures: render_slide_06(),
    7: render_slide_07,
    8: render_slide_08,
    9: render_slide_09,
    10: render_slide_10,
    11: lambda pictures: render_slide_11(),
    12: lambda pictures: render_slide_12(),
    13: render_slide_13,
    14: render_slide_14,
    15: render_slide_15,
    16: render_slide_16,
    17: render_slide_17,
    18: render_slide_18,
    19: render_slide_19,
    20: lambda pictures: render_slide_20(),
}


def save_body_image(image: Image.Image, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, format="PNG")


def replace_slide_body(slide, body_path: Path) -> None:
    preserve_ids: set[int] = set()
    shapes = list(slide.shapes)
    for idx, shape in enumerate(shapes):
        if idx < 3:
            preserve_ids.add(id(shape._element))
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP and shape.top < Inches(1.0):
            preserve_ids.add(id(shape._element))
    for shape in list(slide.shapes):
        if id(shape._element) in preserve_ids:
            continue
        remove_shape(shape)
    slide.shapes.add_picture(str(body_path), BODY_BOX[0], BODY_BOX[1], width=BODY_BOX[2], height=BODY_BOX[3])


def clean_cover_slide(slide) -> None:
    for shape in list(slide.shapes):
        if not getattr(shape, "has_text_frame", False):
            continue
        if shape.text.strip() == "HSVmode:HSV":
            remove_shape(shape)


def main() -> None:
    args = parse_args()
    if not args.input.exists():
        raise FileNotFoundError(f"Input PPTX not found: {args.input}")

    prs = Presentation(str(args.input))
    clean_cover_slide(prs.slides[0])

    for slide_idx in range(2, 21):
        slide = prs.slides[slide_idx - 1]
        pictures = extract_picture_shapes(slide)
        renderer = RENDERERS.get(slide_idx)
        if renderer is None:
            continue
        body = renderer(pictures)
        asset_path = args.asset_dir / f"slide_{slide_idx:02d}_body.png"
        save_body_image(body, asset_path)
        replace_slide_body(slide, asset_path)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(args.output))
    print(f"saved_ppt={args.output}")
    print(f"saved_assets_dir={args.asset_dir}")


if __name__ == "__main__":
    main()
