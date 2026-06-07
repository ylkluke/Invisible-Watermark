from __future__ import annotations

import argparse
import json
from pathlib import Path

import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "刘可毕业答辩.pptx"
DEFAULT_OUTPUT = ROOT / "刘可毕业答辩_答辩优化.pptx"
DEFAULT_METRICS = ROOT / "outputs" / "experiment_reports" / "experiment_metrics_summary.json"
DEFAULT_ASSET_DIR = ROOT / "outputs" / "defense_ppt" / "readable_assets"
FONT_NAME = "Microsoft YaHei"

DISPLAY_LABELS = {
    "A0_Full": "A0\nFull",
    "A1_NoRobust": "A1\nNoRob",
    "A2_NoNoiseReg": "A2\nNoNoise",
    "A3_NoHVS": "A3\nNoHVS",
    "A4_NoPayload": "A4\nNoPay",
    "A5_NoSSIMinWM": "A5\nNoSSIM",
}

SLIDE_17_STYLE = {
    "Ablation study": {"size": 13.0, "bold": False},
    "水印 PSNR 均值与 P10-P90 区间": {"size": 17.5, "bold": False},
    "水印 SSIM 均值与 P10-P90 区间": {"size": 17.5, "bold": False},
    "A0_Full": {"size": 14.0, "bold": True},
    "A1_NoRobust": {"size": 14.0, "bold": True},
    "A2_NoNoiseReg": {"size": 14.0, "bold": True},
    "整体最均衡，说明完整模块与损失组合是合理的。": {"size": 13.2, "bold": False},
    "下降最明显，说明鲁棒训练不仅影响攻击场景，也影响整体恢复稳定性。": {"size": 13.2, "bold": False},
    "Attack-all 下 SSIM 下滑更明显，说明噪声正则有助于结构一致性。": {"size": 13.2, "bold": False},
}

SLIDE_18_STYLE = {
    "Ablation follow-up": {"size": 13.0, "bold": False},
    "相对 A0_Full 的 attack 指标变化": {"size": 17.5, "bold": False},
    "各消融变体的宿主图像质量": {"size": 17.5, "bold": False},
    "A3_NoHVS": {"size": 14.0, "bold": True},
    "A4_NoPayload": {"size": 14.0, "bold": True},
    "A5_NoSSIMMinWM": {"size": 14.0, "bold": True},
    "在Clean水印上恢复较高，但并没有带来更好的综合表现。": {"size": 13.2, "bold": False},
    "Host SSIM很高，但水印恢复明显失衡。": {"size": 13.2, "bold": False},
    "结构相似性下降，仍需SSIM监督。": {"size": 13.2, "bold": False},
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate a defense-ready PPT copy with clearer ablation figures and text.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="Source PPTX path.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Optimized PPTX path.")
    parser.add_argument("--metrics", type=Path, default=DEFAULT_METRICS, help="Metrics summary JSON path.")
    parser.add_argument(
        "--asset-dir",
        type=Path,
        default=DEFAULT_ASSET_DIR,
        help="Directory for regenerated readable chart assets.",
    )
    return parser.parse_args()


def load_ablation_rows(metrics_path: Path) -> list[dict[str, float]]:
    data = json.loads(metrics_path.read_text(encoding="utf-8"))
    rows = data["ablation"]["rows"]
    if not rows:
        raise ValueError("No ablation rows found in metrics summary.")
    return rows


def set_plot_style() -> None:
    plt.rcParams["font.family"] = "sans-serif"
    plt.rcParams["font.sans-serif"] = [FONT_NAME, "SimHei", "Noto Sans CJK SC", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False


def _style_axes(ax, *, ylabel: str, title: str, zero_line: bool = False) -> None:
    ax.set_ylabel(ylabel, fontsize=15)
    ax.set_title(title, fontsize=18, weight="bold", pad=12)
    ax.grid(axis="y", linestyle="--", linewidth=0.8, alpha=0.28)
    ax.tick_params(axis="y", labelsize=13)
    ax.tick_params(axis="x", labelsize=12.5, pad=8)
    ax.set_axisbelow(True)
    if zero_line:
        ax.axhline(0, color="#23374d", linewidth=1.3)


def _apply_xticklabels(ax, labels: list[str]) -> None:
    ax.set_xticks(range(len(labels)))
    ax.set_xticklabels(labels)
    for label in ax.get_xticklabels():
        label.set_multialignment("center")
        label.set_linespacing(0.95)


def save_delta_figure(rows: list[dict[str, float]], asset_dir: Path) -> Path:
    base = rows[0]
    labels = [DISPLAY_LABELS[row["variant"]] for row in rows]
    delta_psnr = [row["wm_psnr_attack_all"] - base["wm_psnr_attack_all"] for row in rows]
    delta_ssim = [row["wm_ssim_attack_all"] - base["wm_ssim_attack_all"] for row in rows]
    x = list(range(len(rows)))

    fig, axes = plt.subplots(1, 2, figsize=(14.4, 5.6), dpi=220)
    fig.patch.set_facecolor("white")

    axes[0].bar(x, delta_psnr, color="#5B84B1", width=0.72)
    _apply_xticklabels(axes[0], labels)
    _style_axes(axes[0], ylabel="delta dB", title="Attack WM PSNR delta vs A0", zero_line=True)

    axes[1].bar(x, delta_ssim, color="#FF8C1A", width=0.72)
    _apply_xticklabels(axes[1], labels)
    _style_axes(axes[1], ylabel="delta SSIM", title="Attack WM SSIM delta vs A0", zero_line=True)

    fig.subplots_adjust(left=0.065, right=0.985, top=0.86, bottom=0.26, wspace=0.20)
    out_path = asset_dir / "fig_ablation_delta_vs_a0_attack_readable.png"
    fig.savefig(out_path, facecolor="white")
    plt.close(fig)
    return out_path


def save_host_quality_figure(rows: list[dict[str, float]], asset_dir: Path) -> Path:
    labels = [DISPLAY_LABELS[row["variant"]] for row in rows]
    host_psnr = [row["host_psnr_clean"] for row in rows]
    host_ssim = [row["host_ssim_clean"] for row in rows]
    x = list(range(len(rows)))

    fig, axes = plt.subplots(1, 2, figsize=(14.4, 5.6), dpi=220)
    fig.patch.set_facecolor("white")

    axes[0].bar(x, host_psnr, color="#5AA34A", width=0.72)
    _apply_xticklabels(axes[0], labels)
    _style_axes(axes[0], ylabel="host PSNR (dB)", title="Clean host PSNR")

    axes[1].bar(x, host_ssim, color="#B07AA6", width=0.72)
    _apply_xticklabels(axes[1], labels)
    _style_axes(axes[1], ylabel="host SSIM", title="Clean host SSIM")

    fig.subplots_adjust(left=0.07, right=0.985, top=0.86, bottom=0.26, wspace=0.18)
    out_path = asset_dir / "fig_ablation_host_quality_readable.png"
    fig.savefig(out_path, facecolor="white")
    plt.close(fig)
    return out_path


def remove_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def replace_picture_in_box(slide, old_shape, image_path: Path) -> None:
    left = old_shape.left
    top = old_shape.top
    width = old_shape.width
    height = old_shape.height

    with Image.open(image_path) as img:
        image_aspect = img.width / img.height
    box_aspect = width / height

    if image_aspect >= box_aspect:
        new_width = width
        new_height = int(round(width / image_aspect))
        new_left = left
        new_top = top + (height - new_height) // 2
    else:
        new_height = height
        new_width = int(round(height * image_aspect))
        new_top = top
        new_left = left + (width - new_width) // 2

    remove_shape(old_shape)
    slide.shapes.add_picture(str(image_path), new_left, new_top, width=new_width, height=new_height)


def normalize_text_frame(shape, *, min_size: float | None = None) -> None:
    if not shape.has_text_frame:
        return
    shape.text_frame.word_wrap = True
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            run.font.name = FONT_NAME
            if min_size is not None:
                current = run.font.size.pt if run.font.size else None
                if current is None or current < min_size:
                    run.font.size = Pt(min_size)


def apply_text_style(shape, *, size: float, bold: bool) -> None:
    if not shape.has_text_frame:
        return
    shape.text_frame.word_wrap = True
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            run.font.name = FONT_NAME
            run.font.size = Pt(size)
            run.font.bold = bold


def style_slide_text(slide, exact_styles: dict[str, dict[str, float | bool]], *, min_size: float) -> None:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        normalize_text_frame(shape, min_size=min_size)
        text = shape.text.strip()
        if not text:
            continue
        style = exact_styles.get(text)
        if style:
            apply_text_style(shape, size=float(style["size"]), bold=bool(style["bold"]))


def update_ppt(input_path: Path, output_path: Path, delta_image: Path, host_image: Path) -> None:
    prs = Presentation(str(input_path))

    slide17 = prs.slides[16]
    slide18 = prs.slides[17]

    style_slide_text(slide17, SLIDE_17_STYLE, min_size=13.0)
    style_slide_text(slide18, SLIDE_18_STYLE, min_size=13.0)

    pictures = [shape for shape in slide18.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    if len(pictures) < 2:
        raise ValueError("Slide 18 does not contain the expected ablation figures.")
    pictures.sort(key=lambda shape: shape.left)
    replace_picture_in_box(slide18, pictures[0], delta_image)
    replace_picture_in_box(slide18, pictures[1], host_image)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))


def main() -> None:
    args = parse_args()
    if not args.input.exists():
        raise FileNotFoundError(f"Input PPT not found: {args.input}")
    if not args.metrics.exists():
        raise FileNotFoundError(f"Metrics summary not found: {args.metrics}")

    args.asset_dir.mkdir(parents=True, exist_ok=True)

    set_plot_style()
    rows = load_ablation_rows(args.metrics)
    delta_image = save_delta_figure(rows, args.asset_dir)
    host_image = save_host_quality_figure(rows, args.asset_dir)
    update_ppt(args.input, args.output, delta_image, host_image)

    print(f"saved_ppt={args.output}")
    print(f"saved_asset={delta_image}")
    print(f"saved_asset={host_image}")


if __name__ == "__main__":
    main()
