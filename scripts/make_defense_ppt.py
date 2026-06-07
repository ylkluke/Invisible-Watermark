from __future__ import annotations

import argparse
import json
import mimetypes
import zipfile
from dataclasses import dataclass
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = ROOT / "毕业论文11.docx"
OUTPUT_DIR = ROOT / "outputs" / "defense_ppt"
DOCX_MEDIA_DIR = OUTPUT_DIR / "docx_media"
PPTX_PATH = OUTPUT_DIR / "基于深度学习的数字图像不可见水印算法_毕业答辩.pptx"
MISSING_ASSETS_PATH = OUTPUT_DIR / "缺图清单.md"
SLIDE_SUMMARY_PATH = OUTPUT_DIR / "slides_summary.md"
METRICS_PATH = ROOT / "outputs" / "experiment_reports" / "experiment_metrics_summary.json"
TOTAL_SLIDES = 21
VIDEO_CANDIDATES = [
    OUTPUT_DIR / "streamlit_demo.mp4",
    OUTPUT_DIR / "streamlit_demo.mov",
    OUTPUT_DIR / "streamlit_demo.webm",
    OUTPUT_DIR / "app_demo.mp4",
]
VIDEO_POSTER_CANDIDATES = [
    OUTPUT_DIR / "streamlit_demo_poster.png",
    OUTPUT_DIR / "streamlit_demo_poster.jpg",
]


@dataclass
class MissingAsset:
    slide_title: str
    placeholder: str
    needed_type: str
    suggestion: str


class _NullFill:
    def __init__(self) -> None:
        self.fore_color = self

    def solid(self) -> None:
        return None

    def background(self) -> None:
        return None


class _NullLine:
    def __init__(self) -> None:
        self.color = self
        self.fill = _NullFill()
        self.width = 0


class _NullFont:
    def __init__(self) -> None:
        self.name = ""
        self.size = 0
        self.bold = False
        self.color = self


class _NullRun:
    def __init__(self) -> None:
        self.text = ""
        self.font = _NullFont()


class _NullParagraph:
    def __init__(self) -> None:
        self.text = ""
        self.level = 0
        self.alignment = None
        self.space_after = 0
        self.line_spacing = 0
        self.bullet = False
        self.runs = [_NullRun()]

    def add_run(self) -> _NullRun:
        run = _NullRun()
        self.runs = [run]
        return run


class _NullTextFrame:
    def __init__(self) -> None:
        self.paragraphs = [_NullParagraph()]
        self.word_wrap = True
        self.margin_left = 0
        self.margin_right = 0
        self.margin_top = 0
        self.margin_bottom = 0
        self.vertical_anchor = None

    def clear(self) -> None:
        self.paragraphs = [_NullParagraph()]

    def add_paragraph(self) -> _NullParagraph:
        p = _NullParagraph()
        self.paragraphs.append(p)
        return p


class _NullShape:
    def __init__(self) -> None:
        self.left = 0
        self.top = 0
        self.width = 0
        self.height = 0
        self.fill = _NullFill()
        self.line = _NullLine()
        self.text_frame = _NullTextFrame()


class _NullShapes:
    def __init__(self) -> None:
        self._shape = _NullShape()

    def add_shape(self, *args, **kwargs) -> _NullShape:
        return self._shape

    def add_textbox(self, *args, **kwargs) -> _NullShape:
        return self._shape

    def add_picture(self, *args, **kwargs) -> _NullShape:
        return self._shape

    def add_movie(self, *args, **kwargs) -> _NullShape:
        return self._shape

    def __getitem__(self, idx) -> _NullShape:
        return self._shape


class _NullBackground:
    def __init__(self) -> None:
        self.fill = _NullFill()


class NullSlide:
    def __init__(self) -> None:
        self.is_null = True
        self.shapes = _NullShapes()
        self.background = _NullBackground()


NULL_SLIDE = NullSlide()
NULL_SHAPE = _NullShape()


def rgb(hex_str: str) -> RGBColor:
    hex_str = hex_str.strip().lstrip("#")
    return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))


BG = rgb("F8FBFF")
INK = rgb("102033")
MUTED = rgb("5E6A7D")
ACCENT = rgb("0F4C81")
ACCENT_2 = rgb("1B6CA8")
ACCENT_3 = rgb("5F7FA8")
CARD = rgb("FCFEFF")
LINE = rgb("C9D9EC")
PLACEHOLDER_FILL = rgb("EEF5FF")
PLACEHOLDER_LINE = rgb("A9C0DE")

FONT_HEAD = "Microsoft YaHei"
FONT_BODY = "Microsoft YaHei"
FONT_FORMULA = "Cambria Math"


def ensure_dirs() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    DOCX_MEDIA_DIR.mkdir(parents=True, exist_ok=True)


def extract_docx_media(docx_path: Path, out_dir: Path) -> dict[str, Path]:
    extracted: dict[str, Path] = {}
    with zipfile.ZipFile(docx_path) as zf:
        for name in zf.namelist():
            if not name.startswith("word/media/"):
                continue
            data = zf.read(name)
            target = out_dir / Path(name).name
            target.write_bytes(data)
            extracted[target.name] = target
    return extracted


def load_metrics() -> dict:
    return json.loads(METRICS_PATH.read_text(encoding="utf-8"))


def make_prs() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def is_null_slide(slide) -> bool:
    return bool(getattr(slide, "is_null", False))


def clear_slide(slide) -> None:
    if is_null_slide(slide):
        return
    for shape in list(slide.shapes):
        el = shape.element
        el.getparent().remove(el)


def acquire_slide(prs: Presentation, slide_no: int, target_slides: set[int] | None, update_existing: bool):
    if update_existing:
        if target_slides is not None and slide_no not in target_slides:
            return NULL_SLIDE
        slide = prs.slides[slide_no - 1]
        clear_slide(slide)
        return slide
    return prs.slides.add_slide(prs.slide_layouts[6])


def ensure_slide_capacity(prs: Presentation, total_slides: int) -> None:
    while len(prs.slides) < total_slides:
        prs.slides.add_slide(prs.slide_layouts[6])


def set_bg(slide, color: RGBColor = BG, *, with_header: bool = True) -> None:
    if is_null_slide(slide):
        return
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color
    if with_header:
        add_school_header(slide)


def add_school_header(slide) -> None:
    if is_null_slide(slide):
        return
    add_rect(slide, Inches(10.44), Inches(0.06), Inches(2.40), Inches(0.48), fill_color=rgb("FFFFFF"), line_color=rgb("DDE8F5"))
    slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(10.55), Inches(0.11), Inches(0.34), Inches(0.34)).fill.solid()
    seal = slide.shapes[-1]
    seal.fill.fore_color.rgb = ACCENT
    seal.line.color.rgb = ACCENT
    add_textbox(
        slide,
        Inches(10.55),
        Inches(0.13),
        Inches(0.34),
        Inches(0.18),
        "长大",
        font_size=6.2,
        bold=True,
        color=rgb("FFFFFF"),
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        margin_pt=0,
    )
    add_textbox(
        slide,
        Inches(10.98),
        Inches(0.10),
        Inches(1.78),
        Inches(0.20),
        "长安大学",
        font_size=10.5,
        bold=True,
        color=INK,
        align=PP_ALIGN.RIGHT,
        valign=MSO_ANCHOR.MIDDLE,
        margin_pt=0,
    )
    add_textbox(
        slide,
        Inches(10.98),
        Inches(0.31),
        Inches(1.78),
        Inches(0.14),
        "CHANG'AN UNIVERSITY",
        font_size=5.8,
        bold=True,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
        valign=MSO_ANCHOR.MIDDLE,
        margin_pt=0,
    )
    add_rect(slide, Inches(10.98), Inches(0.48), Inches(1.78), Inches(0.015), fill_color=rgb("DDE8F5"), line_color=rgb("DDE8F5"), radius=False)


def add_rect(slide, x, y, w, h, *, fill_color=CARD, line_color=LINE, radius=True):
    if is_null_slide(slide):
        return NULL_SHAPE
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, x, y, w, h)
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = fill_color
    line = shape.line
    line.color.rgb = line_color
    line.width = Pt(1.0)
    return shape


def add_textbox(
    slide,
    x,
    y,
    w,
    h,
    text: str,
    *,
    font_size: float = 20,
    bold: bool = False,
    color: RGBColor = INK,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin_pt: float = 8,
    font_name: str = FONT_BODY,
):
    if is_null_slide(slide):
        return NULL_SHAPE
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Pt(margin_pt)
    tf.margin_right = Pt(margin_pt)
    tf.margin_top = Pt(margin_pt)
    tf.margin_bottom = Pt(margin_pt)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    if is_null_slide(slide):
        return
    add_textbox(
        slide,
        Inches(0.55),
        Inches(0.28),
        Inches(9.8),
        Inches(0.55),
        title,
        font_size=26,
        bold=True,
        color=INK,
        font_name=FONT_HEAD,
    )
    slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(0.82), Inches(1.25), Inches(0.06)).fill.solid()
    bar = slide.shapes[-1]
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    if subtitle:
        add_textbox(
            slide,
            Inches(1.95),
            Inches(0.67),
            Inches(4.0),
            Inches(0.3),
            subtitle,
            font_size=10.5,
            color=MUTED,
            font_name=FONT_BODY,
        )


def add_bullets(slide, x, y, w, h, items: list[str], *, font_size: float = 19, color: RGBColor = INK) -> None:
    if is_null_slide(slide):
        return
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Pt(4)
    tf.margin_right = Pt(4)
    tf.margin_top = Pt(4)
    tf.margin_bottom = Pt(4)
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)
        p.line_spacing = 1.18
        p.bullet = True
        if p.runs:
            run = p.runs[0]
            run.font.name = FONT_BODY
            run.font.size = Pt(font_size)
            run.font.color.rgb = color


def add_card(slide, x, y, w, h, title: str, body: str, *, accent: RGBColor = ACCENT) -> None:
    if is_null_slide(slide):
        return
    card = add_rect(slide, x, y, w, h)
    slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, Inches(0.08), h).fill.solid()
    pill = slide.shapes[-1]
    pill.fill.fore_color.rgb = accent
    pill.line.fill.background()
    add_textbox(slide, x + Inches(0.12), y + Inches(0.06), w - Inches(0.18), Inches(0.24), title, font_size=12, bold=True)
    add_textbox(slide, x + Inches(0.12), y + Inches(0.31), w - Inches(0.18), h - Inches(0.35), body, font_size=11.5, color=MUTED)


def add_formula_card(
    slide,
    x,
    y,
    w,
    h,
    label: str,
    formula: str,
    note: str,
    *,
    accent: RGBColor = ACCENT,
    formula_size: float = 16,
    note_size: float = 9.8,
    formula_align=PP_ALIGN.CENTER,
):
    if is_null_slide(slide):
        return
    add_rect(slide, x, y, w, h, fill_color=CARD, line_color=LINE)
    add_rect(slide, x + Inches(0.12), y + Inches(0.12), Inches(0.92), Inches(0.28), fill_color=rgb("EAF2FF"), line_color=rgb("EAF2FF"))
    add_textbox(
        slide,
        x + Inches(0.12),
        y + Inches(0.12),
        Inches(0.92),
        Inches(0.28),
        label,
        font_size=9.8,
        bold=True,
        color=accent,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        margin_pt=2,
    )
    add_textbox(
        slide,
        x + Inches(0.16),
        y + Inches(0.48),
        w - Inches(0.32),
        h - Inches(1.0),
        formula,
        font_size=formula_size,
        bold=True,
        color=INK,
        align=formula_align,
        valign=MSO_ANCHOR.MIDDLE,
        margin_pt=4,
        font_name=FONT_FORMULA,
    )
    add_rect(slide, x + Inches(0.16), y + h - Inches(0.42), w - Inches(0.32), Inches(0.02), fill_color=accent, line_color=accent, radius=False)
    add_textbox(
        slide,
        x + Inches(0.16),
        y + h - Inches(0.34),
        w - Inches(0.32),
        Inches(0.18),
        note,
        font_size=note_size,
        color=MUTED,
        align=PP_ALIGN.LEFT,
        valign=MSO_ANCHOR.MIDDLE,
        margin_pt=2,
    )


def add_metric_card(slide, x, y, w, h, label: str, value: str, note: str) -> None:
    if is_null_slide(slide):
        return
    add_rect(slide, x, y, w, h, fill_color=CARD, line_color=LINE)
    add_textbox(slide, x + Inches(0.08), y + Inches(0.05), w - Inches(0.16), Inches(0.2), label, font_size=10.5, bold=True, color=MUTED)
    add_textbox(slide, x + Inches(0.08), y + Inches(0.22), w - Inches(0.16), Inches(0.28), value, font_size=18, bold=True, color=ACCENT)
    add_textbox(slide, x + Inches(0.08), y + Inches(0.47), w - Inches(0.16), Inches(0.18), note, font_size=9.5, color=MUTED)


def image_size(path: Path) -> tuple[int, int]:
    with Image.open(path) as im:
        return im.size


def first_existing(paths: list[Path]) -> Path | None:
    for path in paths:
        if path.exists():
            return path
    return None


def guess_mime_type(path: Path) -> str:
    mime_type, _ = mimetypes.guess_type(str(path))
    return mime_type or "video/unknown"


def add_picture_fit(slide, path: Path, x, y, w, h, *, crop: bool = False):
    if is_null_slide(slide):
        return
    iw, ih = image_size(path)
    frame_ratio = w / h
    img_ratio = iw / ih
    if crop:
        if img_ratio >= frame_ratio:
            pic_h = h
            pic_w = h * img_ratio
            pic_x = x - (pic_w - w) / 2
            pic_y = y
        else:
            pic_w = w
            pic_h = w / img_ratio
            pic_x = x
            pic_y = y - (pic_h - h) / 2
    else:
        if img_ratio >= frame_ratio:
            pic_w = w
            pic_h = w / img_ratio
            pic_x = x
            pic_y = y + (h - pic_h) / 2
        else:
            pic_h = h
            pic_w = h * img_ratio
            pic_x = x + (w - pic_w) / 2
            pic_y = y
    slide.shapes.add_picture(str(path), pic_x, pic_y, width=pic_w, height=pic_h)


def add_picture_or_placeholder(
    slide,
    *,
    path: Path | None,
    x,
    y,
    w,
    h,
    caption: str | None,
    slide_title: str,
    placeholder_name: str,
    needed_type: str,
    suggestion: str,
    missing_assets: list[MissingAsset],
) -> None:
    if is_null_slide(slide):
        return
    add_rect(slide, x, y, w, h, fill_color=CARD, line_color=LINE, radius=False)
    if path is not None and path.exists():
        add_picture_fit(slide, path, x + Inches(0.02), y + Inches(0.02), w - Inches(0.04), h - Inches(0.04), crop=False)
    else:
        add_rect(slide, x + Inches(0.06), y + Inches(0.06), w - Inches(0.12), h - Inches(0.12), fill_color=PLACEHOLDER_FILL, line_color=PLACEHOLDER_LINE)
        missing_assets.append(MissingAsset(slide_title, placeholder_name, needed_type, suggestion))
    if caption:
        add_textbox(slide, x, y + h + Inches(0.03), w, Inches(0.22), caption, font_size=9.5, color=MUTED, align=PP_ALIGN.CENTER)


def add_video_or_placeholder(
    slide,
    *,
    video_path: Path | None,
    poster_path: Path | None,
    x,
    y,
    w,
    h,
    caption: str | None,
    slide_title: str,
    placeholder_name: str,
    needed_type: str,
    suggestion: str,
    missing_assets: list[MissingAsset],
) -> None:
    if is_null_slide(slide):
        return
    add_rect(slide, x, y, w, h, fill_color=CARD, line_color=LINE, radius=False)
    if video_path is not None and video_path.exists():
        slide.shapes.add_movie(
            str(video_path),
            x + Inches(0.03),
            y + Inches(0.03),
            w - Inches(0.06),
            h - Inches(0.06),
            poster_frame_image=str(poster_path) if poster_path is not None and poster_path.exists() else None,
            mime_type=guess_mime_type(video_path),
        )
    else:
        add_rect(slide, x + Inches(0.06), y + Inches(0.06), w - Inches(0.12), h - Inches(0.12), fill_color=PLACEHOLDER_FILL, line_color=PLACEHOLDER_LINE)
        missing_assets.append(MissingAsset(slide_title, placeholder_name, needed_type, suggestion))
    if caption:
        add_textbox(slide, x, y + h + Inches(0.03), w, Inches(0.22), caption, font_size=9.5, color=MUTED, align=PP_ALIGN.CENTER)


def set_notes(slide, script: str) -> None:
    if is_null_slide(slide):
        return
    tf = slide.notes_slide.notes_text_frame
    tf.clear()
    paragraphs = [part.strip() for part in script.strip().split("\n\n") if part.strip()]
    if not paragraphs:
        return
    first = tf.paragraphs[0]
    first.text = paragraphs[0]
    for part in paragraphs[1:]:
        p = tf.add_paragraph()
        p.text = part


def add_chip(slide, x, y, w, h, text: str, *, fill_color: RGBColor = rgb("EAF2FF"), text_color: RGBColor = ACCENT) -> None:
    if is_null_slide(slide):
        return
    add_rect(slide, x, y, w, h, fill_color=fill_color, line_color=fill_color)
    add_textbox(slide, x, y, w, h, text, font_size=10, bold=True, color=text_color, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin_pt=2)


def add_kv_item(slide, x, y, w, h, title: str, body: str, *, accent: RGBColor = ACCENT) -> None:
    if is_null_slide(slide):
        return
    add_rect(slide, x, y, w, h, fill_color=CARD, line_color=LINE)
    add_rect(slide, x, y, Inches(0.07), h, fill_color=accent, line_color=accent, radius=False)
    add_textbox(slide, x + Inches(0.16), y + Inches(0.08), w - Inches(0.24), Inches(0.22), title, font_size=11.5, bold=True, color=INK)
    add_textbox(slide, x + Inches(0.16), y + Inches(0.34), w - Inches(0.24), h - Inches(0.42), body, font_size=10.8, color=MUTED)


def build_presentation(
    prs: Presentation,
    media: dict[str, Path],
    metrics: dict,
    missing_assets: list[MissingAsset],
    *,
    target_slides: set[int] | None = None,
    update_existing: bool = False,
) -> list[tuple[str, list[str]]]:
    slide_summaries: list[tuple[str, list[str]]] = []

    fig_dir = ROOT / "outputs" / "experiment_reports" / "figures"
    demo_dir = ROOT / "outputs" / "edge_wm_noise_test" / "compare_r1_20260411_robust_attack_all" / "samples" / "微信图片_20260301215540_18_16"
    demo_video = first_existing(VIDEO_CANDIDATES)

    def media_path(name: str) -> Path | None:
        return media.get(name)

    def fig_path(name: str) -> Path:
        return fig_dir / name

    def demo_path(name: str) -> Path:
        return demo_dir / name

    demo_video_poster = first_existing(VIDEO_POSTER_CANDIDATES + [demo_path("cover.png"), demo_path("container.png")])

    main_clean = metrics["main"]["clean"]
    main_attack = metrics["main"]["attack_all"]
    compare = metrics["compare"]

    # 1. Cover
    slide = acquire_slide(prs, 1, target_slides, update_existing)
    set_bg(slide, rgb("EFF4FB"), with_header=False)
    accent_band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.0), Inches(0.0), Inches(13.333), Inches(0.28))
    accent_band.fill.solid()
    accent_band.fill.fore_color.rgb = ACCENT
    accent_band.line.fill.background()
    add_school_header(slide)
    add_textbox(slide, Inches(0.75), Inches(1.00), Inches(7.75), Inches(0.88), "基于深度学习的数字图像不可见水印算法", font_size=27, bold=True, color=INK, font_name=FONT_HEAD)
    add_textbox(slide, Inches(0.78), Inches(1.98), Inches(7.2), Inches(0.36), "本科毕业设计答辩", font_size=13, color=MUTED)
    add_chip(slide, Inches(0.78), Inches(2.45), Inches(1.45), Inches(0.34), "Y通道载荷")
    add_chip(slide, Inches(2.35), Inches(2.45), Inches(1.35), Inches(0.34), "RGB残差嵌入", fill_color=rgb("EAFBF7"), text_color=ACCENT_2)
    add_chip(slide, Inches(3.82), Inches(2.45), Inches(1.30), Inches(0.34), "鲁棒训练", fill_color=rgb("FFF4E0"), text_color=ACCENT_3)
    info = add_rect(slide, Inches(0.78), Inches(3.15), Inches(4.35), Inches(1.55), fill_color=CARD, line_color=LINE)
    add_textbox(slide, info.left + Inches(0.1), info.top + Inches(0.12), info.width - Inches(0.2), Inches(0.95), "学生：刘可\n学院：信息工程学院\n专业：电子信息工程\n指导教师：张晓博", font_size=15.5, color=INK)
    panel = add_rect(slide, Inches(8.72), Inches(0.98), Inches(3.62), Inches(4.05), fill_color=rgb("F6FAFF"), line_color=LINE)
    add_textbox(slide, panel.left + Inches(0.18), panel.top + Inches(0.2), panel.width - Inches(0.36), Inches(0.28), "Defense Presentation", font_size=12.5, bold=True, color=ACCENT)
    add_textbox(slide, panel.left + Inches(0.18), panel.top + Inches(0.56), panel.width - Inches(0.36), Inches(0.62), "不可见水印\n深度学习与鲁棒恢复", font_size=18, bold=True, color=INK)
    add_textbox(slide, panel.left + Inches(0.18), panel.top + Inches(1.35), panel.width - Inches(0.36), Inches(0.42), "方法设计 / 实验验证 / 系统演示", font_size=12.5, color=MUTED)
    add_rect(slide, panel.left + Inches(0.18), panel.top + Inches(2.0), Inches(2.9), Inches(0.08), fill_color=ACCENT, line_color=ACCENT, radius=False)
    add_rect(slide, panel.left + Inches(0.18), panel.top + Inches(2.34), Inches(2.2), Inches(0.05), fill_color=ACCENT_2, line_color=ACCENT_2, radius=False)
    add_rect(slide, panel.left + Inches(0.18), panel.top + Inches(2.64), Inches(1.55), Inches(0.05), fill_color=ACCENT_3, line_color=ACCENT_3, radius=False)
    add_chip(slide, panel.left + Inches(0.18), panel.top + Inches(3.08), Inches(1.05), Inches(0.34), "PyTorch", fill_color=rgb("EAF2FF"), text_color=ACCENT)
    add_chip(slide, panel.left + Inches(1.37), panel.top + Inches(3.08), Inches(1.2), Inches(0.34), "Robust", fill_color=rgb("E9F8FF"), text_color=ACCENT_2)
    add_chip(slide, panel.left + Inches(2.71), panel.top + Inches(3.08), Inches(0.7), Inches(0.34), "App", fill_color=rgb("EEF3FF"), text_color=ACCENT_3)
    set_notes(slide, "各位老师好，我是信息工程学院电子信息工程专业的刘可。我的毕业设计题目是《基于深度学习的数字图像不可见水印算法》，指导教师是张晓博老师。\n\n本次汇报主要围绕三个问题展开：第一，为什么要研究不可见数字水印；第二，本文如何通过 Y 通道载荷和残差嵌入实现水印写入与恢复；第三，实验结果能否说明方法在不可见性、恢复质量和鲁棒性之间取得了较好的平衡。")
    slide_summaries.append(("封面", ["课题题目", "学生信息", "工程风视觉面板"]))

    # 2. Agenda
    slide = acquire_slide(prs, 2, target_slides, update_existing)
    set_bg(slide, rgb("F4F9FF"))
    add_title(slide, "目录", "Contents")
    add_textbox(slide, Inches(0.78), Inches(1.18), Inches(4.1), Inches(0.4), "汇报结构", font_size=18, bold=True, color=ACCENT)
    agenda_items = [
        ("01", "研究背景与问题定义", "图像传播场景、版权保护需求与本文目标。"),
        ("02", "方法设计与模型结构", "Y 通道载荷、残差嵌入、恢复链路与损失函数。"),
        ("03", "实验结果与对比分析", "主实验、Plain vs Robust 对比实验和消融实验。"),
        ("04", "系统演示", "基于 Streamlit 的交互式水印嵌入与恢复演示。"),
        ("05", "结论与展望", "本文结论、当前不足和后续改进方向。"),
    ]
    for idx, (code, title, body) in enumerate(agenda_items):
        x = Inches(0.78) if idx < 3 else Inches(6.7)
        y = Inches(1.78 + (idx % 3) * 1.55)
        add_rect(slide, x, y, Inches(5.0), Inches(1.22), fill_color=CARD, line_color=LINE)
        add_textbox(slide, x + Inches(0.14), y + Inches(0.16), Inches(0.7), Inches(0.34), code, font_size=18, bold=True, color=ACCENT)
        add_textbox(slide, x + Inches(0.9), y + Inches(0.13), Inches(3.65), Inches(0.28), title, font_size=15, bold=True, color=INK)
        add_textbox(slide, x + Inches(0.9), y + Inches(0.48), Inches(3.8), Inches(0.42), body, font_size=11.5, color=MUTED)
        add_rect(slide, x + Inches(4.58), y + Inches(0.16), Inches(0.2), Inches(0.9), fill_color=ACCENT if idx % 2 == 0 else ACCENT_2, line_color=LINE, radius=False)
    set_notes(slide, "我的汇报分为五个部分。首先介绍研究背景和问题定义，其次说明本文的模型结构和关键模块，然后给出实验设置、主实验结果、对比实验和消融分析，接着展示基于 Streamlit 的交互式演示系统，最后总结本文工作并说明后续可以继续改进的方向。")
    slide_summaries.append(("目录", ["五部分汇报结构", "蓝白工程风目录卡片"]))

    # 3. Background
    slide = acquire_slide(prs, 3, target_slides, update_existing)
    set_bg(slide)
    add_title(slide, "研究背景与目标", "Why this problem matters")
    add_card(slide, Inches(0.72), Inches(1.22), Inches(3.85), Inches(1.1), "应用场景", "社交媒体、在线图库、AIGC 内容传播链路中，图像会被频繁复制、压缩、裁剪和二次编辑。", accent=ACCENT)
    add_card(slide, Inches(4.74), Inches(1.22), Inches(3.85), Inches(1.1), "核心矛盾", "不可见性、可恢复性和鲁棒性难以同时兼顾，传统规则式方法在复杂攻击下容易退化。", accent=ACCENT_2)
    add_card(slide, Inches(8.76), Inches(1.22), Inches(3.85), Inches(1.1), "论文目标", "在保证宿主图像视觉质量的前提下，让水印可以稳定恢复，并对常见攻击保持一定鲁棒性。", accent=ACCENT_3)
    bullets = [
        "数字图像版权保护需要“平时看不见、需要时能取出”的技术手段。",
        "真实传播链路不只包含单一压缩，还会叠加裁剪、缩放、模糊和马赛克等失真。",
        "论文不只做模型设计，也补齐了训练、推理、批测、对比实验和消融实验的工程链路。",
    ]
    add_rect(slide, Inches(0.72), Inches(2.62), Inches(5.55), Inches(3.85))
    add_bullets(slide, Inches(0.9), Inches(2.85), Inches(5.15), Inches(3.4), bullets, font_size=18)
    add_rect(slide, Inches(6.55), Inches(2.62), Inches(6.05), Inches(3.85))
    add_textbox(slide, Inches(6.78), Inches(2.9), Inches(5.6), Inches(0.3), "核心价值", font_size=14, bold=True, color=MUTED)
    add_card(slide, Inches(6.82), Inches(3.35), Inches(1.7), Inches(2.15), "确权", "嵌入版权标识，服务作者身份确认。", accent=ACCENT)
    add_card(slide, Inches(8.72), Inches(3.35), Inches(1.7), Inches(2.15), "溯源", "图像被传播后仍能回溯来源。", accent=ACCENT_2)
    add_card(slide, Inches(10.62), Inches(3.35), Inches(1.7), Inches(2.15), "复现", "代码与指标链路完整，方便论文验证。", accent=ACCENT_3)
    set_notes(slide, "首先看研究背景。随着互联网平台和 AIGC 内容生产的发展，数字图像的复制、转发和二次编辑成本越来越低，版权归属不清、内容被盗用以及传播后难以溯源的问题也更加突出。\n\n传统可见水印虽然直观，但会影响图像观感；不可见水印更适合真实应用，但需要同时兼顾不可见性、可恢复性和鲁棒性。本文的目标就是在尽量不影响宿主图像视觉质量的前提下，把版权信息稳定地嵌入图像，并在常见攻击后尽可能恢复出来。")
    slide_summaries.append(("研究背景与目标", ["应用场景", "核心矛盾", "论文目标"]))

    # 4. Problem and core idea
    slide = acquire_slide(prs, 4, target_slides, update_existing)
    set_bg(slide)
    add_title(slide, "研究问题与核心思路", "Problem definition")
    add_rect(slide, Inches(0.72), Inches(1.18), Inches(4.35), Inches(5.95))
    add_textbox(slide, Inches(0.92), Inches(1.4), Inches(3.95), Inches(0.3), "核心问题", font_size=16, bold=True, color=ACCENT)
    add_kv_item(slide, Inches(0.92), Inches(1.86), Inches(3.95), Inches(1.02), "问题一", "在保证视觉质量的前提下，把版权信息隐藏进宿主图像。", accent=ACCENT)
    add_kv_item(slide, Inches(0.92), Inches(3.02), Inches(3.95), Inches(1.02), "问题二", "面对裁剪、JPEG、缩放、模糊和马赛克等攻击，恢复主要水印结构。", accent=ACCENT_2)
    add_kv_item(slide, Inches(0.92), Inches(4.18), Inches(3.95), Inches(1.02), "问题三", "把训练、推理、批量测试、对比实验和消融实验串成可复现流程。", accent=ACCENT_3)
    add_rect(slide, Inches(0.92), Inches(5.46), Inches(3.95), Inches(1.18), fill_color=rgb("EEF5FF"), line_color=LINE)
    add_textbox(slide, Inches(1.08), Inches(5.63), Inches(3.62), Inches(0.82), "核心思路：只编码水印模板的 Y 通道亮度结构，将其压缩成单通道噪声载荷，再通过 RGB 残差嵌入网络写入宿主图像。", font_size=12.8, color=INK)
    add_picture_or_placeholder(
        slide,
        path=media_path("image3.png"),
        x=Inches(5.42),
        y=Inches(1.38),
        w=Inches(6.9),
        h=Inches(4.52),
        caption="深度学习水印基本框架示意图",
        slide_title="研究问题与核心思路",
        placeholder_name="深度学习水印流程图",
        needed_type="编码器-攻击层-解码器流程示意图",
        suggestion="若有更清晰矢量版，可替换当前 Word 导出的 PNG",
        missing_assets=missing_assets,
    )
    add_rect(slide, Inches(5.42), Inches(6.02), Inches(6.9), Inches(0.62), fill_color=CARD, line_color=LINE)
    add_textbox(slide, Inches(5.6), Inches(6.18), Inches(6.5), Inches(0.22), "流程图说明的是通用编码器-攻击层-解码器框架，本文在此基础上进一步落到 Y 通道载荷与 RGB 残差嵌入。", font_size=11.4, color=MUTED)
    set_notes(slide, "围绕这个目标，本文主要处理三个问题。第一，水印要能够嵌入宿主图像，但不能产生明显视觉痕迹；第二，图像经过裁剪、压缩、缩放、模糊或马赛克等处理后，仍然要尽量恢复出水印的主体结构；第三，整个方法需要有可复现的训练、推理和测试流程。\n\n本文的核心思路是只编码水印模板的 Y 通道亮度结构，把它压缩为单通道噪声载荷，再通过 RGB 残差嵌入网络写入宿主图像。这样可以把水印信息写入得更隐蔽，也便于后续恢复。")
    slide_summaries.append(("研究问题与核心思路", ["三项研究问题", "Y通道载荷思路", "基础流程图"]))

    # 5. Overall flowchart
    slide = acquire_slide(prs, 5, target_slides, update_existing)
    set_bg(slide)
    add_title(slide, "总体方案流程图", "End-to-end pipeline")
    add_rect(slide, Inches(0.78), Inches(1.16), Inches(11.82), Inches(4.98), fill_color=CARD, line_color=LINE, radius=False)
    add_picture_or_placeholder(
        slide,
        path=media_path("image5.png"),
        x=Inches(0.92),
        y=Inches(1.34),
        w=Inches(11.54),
        h=Inches(4.6),
        caption="论文第 4 章总体结构图",
        slide_title="总体方案流程图",
        placeholder_name="总体结构图",
        needed_type="模型整体架构图",
        suggestion="如有原始 SVG/PDF，可替换当前 PNG 以增强放映清晰度",
        missing_assets=missing_assets,
    )
    add_card(slide, Inches(0.82), Inches(6.24), Inches(3.72), Inches(0.78), "输入侧", "宿主图像走 RGB，水印模板先转 YCbCr，并只保留 Y 通道参与学习。", accent=ACCENT)
    add_card(slide, Inches(4.80), Inches(6.24), Inches(3.72), Inches(0.78), "嵌入侧", "编码后的 payload 与 cover 一起进入 EmbedUNet，预测低幅值 RGB residual。", accent=ACCENT_2)
    add_card(slide, Inches(8.78), Inches(6.24), Inches(3.72), Inches(0.78), "恢复侧", "由 container 先恢复 payload，再解码回 watermark Y，完成版权信息读取。", accent=ACCENT_3)
    set_notes(slide, "这里展示的是本文方法的整体流程。输入包括宿主图像和水印模板。水印模板先转换到 YCbCr 空间，只取其中的 Y 通道参与训练；随后通过编码网络生成单通道 payload。\n\n嵌入端把 payload 和宿主图像一起送入残差嵌入网络，预测一个幅值较小的 RGB 残差，并与原图相加得到含水印图像。恢复端则先从含水印图像中恢复 payload，再由解码网络重建 watermark Y。")
    slide_summaries.append(("总体方案流程图", ["单页单流程图", "输入/嵌入/恢复三段解读"]))

    # 6. Core formulas
    slide = acquire_slide(prs, 6, target_slides, update_existing)
    set_bg(slide)
    add_title(slide, "核心公式：嵌入与恢复", "Method equations")
    add_rect(slide, Inches(0.78), Inches(1.16), Inches(11.82), Inches(0.72), fill_color=rgb("EEF5FF"), line_color=LINE)
    add_textbox(slide, Inches(0.98), Inches(1.34), Inches(11.4), Inches(0.26), "方法链路由 Y 通道载荷、残差嵌入和两阶段恢复组成，下面用公式概括核心计算过程。", font_size=12.4, color=INK)
    add_formula_card(
        slide,
        Inches(0.78),
        Inches(2.06),
        Inches(5.70),
        Inches(1.42),
        "输入表示",
        "W_Y = Y(W_rgb)",
        "先抽取水印亮度结构，降低需要写入和恢复的信息复杂度。",
        accent=ACCENT,
        formula_size=19,
    )
    add_formula_card(
        slide,
        Inches(6.90),
        Inches(2.06),
        Inches(5.70),
        Inches(1.42),
        "载荷编码",
        "P = f_enc(W_Y)",
        "把 watermark Y 映射到单通道噪声载荷 payload。",
        accent=ACCENT_2,
        formula_size=19,
    )
    add_formula_card(
        slide,
        Inches(0.78),
        Inches(3.76),
        Inches(5.70),
        Inches(1.42),
        "残差嵌入",
        "R = alpha * tanh(f_emb(C, P))",
        "通过受限残差写入版权信息，核心目标是减小宿主图像可感知失真。",
        accent=ACCENT_3,
        formula_size=15.6,
    )
    add_formula_card(
        slide,
        Inches(6.90),
        Inches(3.76),
        Inches(5.70),
        Inches(1.42),
        "含水印图生成",
        "C_w = clip(C + R, 0, 1)",
        "残差与宿主图像相加后裁剪到合法像素范围。",
        accent=ACCENT,
        formula_size=16.8,
    )
    add_formula_card(
        slide,
        Inches(0.78),
        Inches(5.46),
        Inches(5.90),
        Inches(1.12),
        "恢复解码",
        "W_hat_Y = f_dec(f_rev(C_w))",
        "先从 container 中恢复 payload，再解码为 watermark Y。",
        accent=ACCENT_2,
        formula_size=15.6,
    )
    add_rect(slide, Inches(6.98), Inches(5.46), Inches(5.62), Inches(1.12), fill_color=CARD, line_color=LINE)
    add_textbox(slide, Inches(7.18), Inches(5.64), Inches(5.18), Inches(0.2), "链路总览", font_size=12.2, bold=True, color=ACCENT)
    add_textbox(slide, Inches(7.18), Inches(5.92), Inches(5.18), Inches(0.44), "Watermark Y -> Payload -> Residual Embed -> Container -> Reveal -> Recover Y", font_size=12.8, bold=True, color=INK)
    add_textbox(slide, Inches(7.18), Inches(6.30), Inches(5.18), Inches(0.16), "方法表达重点是“Y 通道压缩 + 低幅值残差写入 + 两阶段恢复”。", font_size=9.8, color=MUTED)
    set_notes(slide, "方法主线可以用这五个公式概括。首先，彩色水印只取 Y 通道，记为 W_Y；然后通过编码器得到 payload。接着，嵌入网络根据宿主图像 C 和 payload 预测残差 R，并通过 tanh 和系数 alpha 限制残差幅值。\n\n生成含水印图像时，将宿主图像和残差相加，再裁剪到合法像素范围。恢复阶段先从含水印图像中恢复 payload，再解码得到 W_hat_Y。也就是说，本文不是直接把水印图案叠到图像上，而是通过受控残差把水印信息隐式写入。")
    slide_summaries.append(("核心公式：嵌入与恢复", ["五个核心公式卡", "方法链路总览"]))

    # 7. Watermark codec module
    slide = acquire_slide(prs, 7, target_slides, update_existing)
    set_bg(slide)
    add_title(slide, "关键模块：水印编码与解码", "Watermark codec")
    add_rect(slide, Inches(0.78), Inches(1.16), Inches(7.58), Inches(5.28), fill_color=CARD, line_color=LINE, radius=False)
    add_picture_or_placeholder(
        slide,
        path=media_path("image7.png"),
        x=Inches(0.98),
        y=Inches(1.42),
        w=Inches(7.18),
        h=Inches(4.68),
        caption="水印 Y 通道与噪声载荷 payload 之间的编码/解码结构",
        slide_title="关键模块：水印编码与解码",
        placeholder_name="水印编码与解码",
        needed_type="水印编码与解码模块结构图",
        suggestion="若有高分辨率矢量版，优先替换 Word 导出的图片",
        missing_assets=missing_assets,
    )
    add_rect(slide, Inches(8.66), Inches(1.16), Inches(3.82), Inches(5.28))
    add_textbox(slide, Inches(8.88), Inches(1.40), Inches(3.36), Inches(0.30), "模块作用", font_size=16, bold=True, color=ACCENT)
    add_formula_card(slide, Inches(8.88), Inches(1.94), Inches(3.36), Inches(1.04), "编码", "P = f_enc(W_Y)", "将 watermark Y 映射为单通道 payload。", accent=ACCENT, formula_size=16)
    add_formula_card(slide, Inches(8.88), Inches(3.18), Inches(3.36), Inches(1.04), "解码", "W_hat_Y = f_dec(P_hat)", "由恢复出的 payload 重建水印亮度结构。", accent=ACCENT_2, formula_size=15)
    add_card(slide, Inches(8.88), Inches(4.52), Inches(3.36), Inches(1.18), "设计重点", "只恢复水印的 Y 通道，减少颜色信息带来的额外难度，让模型优先学习稳定的亮度结构。", accent=ACCENT_3)
    set_notes(slide, "第一个关键模块是水印编码与解码模块。它解决的是水印如何表示的问题。本文没有直接学习完整彩色水印，而是先提取水印的 Y 通道亮度结构，再把这个结构编码成单通道 payload。\n\n恢复时，模型先估计 payload，再通过解码网络恢复 watermark Y。这样做的好处是降低了需要恢复的信息量，同时让网络更集中地学习水印的主体轮廓和亮度结构。")
    slide_summaries.append(("关键模块：水印编码与解码", ["单模块大图", "编码/解码公式", "Y 通道设计说明"]))

    # 8. Embedder module
    slide = acquire_slide(prs, 8, target_slides, update_existing)
    set_bg(slide)
    add_title(slide, "关键模块：EmbedUNet 残差嵌入", "Residual embedding")
    add_rect(slide, Inches(0.78), Inches(1.16), Inches(7.58), Inches(5.28), fill_color=CARD, line_color=LINE, radius=False)
    add_picture_or_placeholder(
        slide,
        path=media_path("image9.png"),
        x=Inches(0.98),
        y=Inches(1.42),
        w=Inches(7.18),
        h=Inches(4.68),
        caption="EmbedUNet 结合宿主图像和 payload，预测低幅值 RGB residual",
        slide_title="关键模块：EmbedUNet 残差嵌入",
        placeholder_name="EmbedUNet 残差嵌入",
        needed_type="EmbedUNet 残差嵌入模块结构图",
        suggestion="若有高分辨率矢量版，优先替换 Word 导出的图片",
        missing_assets=missing_assets,
    )
    add_rect(slide, Inches(8.66), Inches(1.16), Inches(3.82), Inches(5.28))
    add_textbox(slide, Inches(8.88), Inches(1.40), Inches(3.36), Inches(0.30), "模块作用", font_size=16, bold=True, color=ACCENT)
    add_formula_card(slide, Inches(8.88), Inches(1.94), Inches(3.36), Inches(1.16), "输入", "concat(C, P)", "宿主图像和 payload 在通道维度联合输入。", accent=ACCENT, formula_size=16)
    add_formula_card(slide, Inches(8.88), Inches(3.26), Inches(3.36), Inches(1.16), "输出", "R = alpha * tanh(f_emb(C, P))", "输出受限 RGB residual，避免明显破坏宿主图像。", accent=ACCENT_2, formula_size=12.3)
    add_card(slide, Inches(8.88), Inches(4.72), Inches(3.36), Inches(1.00), "设计重点", "嵌入阶段写入的是残差而不是直接重绘图像，因此更容易控制不可见性。", accent=ACCENT_3)
    set_notes(slide, "第二个关键模块是 EmbedUNet 残差嵌入模块。它的输入是宿主图像和前面得到的 payload，输出不是一张新的完整图像，而是一个 RGB residual。\n\n这里采用残差嵌入，是因为不可见水印最核心的要求是控制图像改变量。模型只需要学习在原图基础上添加怎样的微弱扰动，就能把水印信息写进去。最终的 container 是 cover 加 residual 得到的，因此宿主图像的整体视觉内容可以保持稳定。")
    slide_summaries.append(("关键模块：EmbedUNet 残差嵌入", ["单模块大图", "输入/输出公式", "残差嵌入说明"]))

    # 9. Reveal payload module
    slide = acquire_slide(prs, 9, target_slides, update_existing)
    set_bg(slide)
    add_title(slide, "关键模块：载荷恢复模块", "Payload reveal")
    add_rect(slide, Inches(0.78), Inches(1.16), Inches(7.58), Inches(5.28), fill_color=CARD, line_color=LINE, radius=False)
    add_picture_or_placeholder(
        slide,
        path=media_path("image11.png"),
        x=Inches(0.98),
        y=Inches(1.42),
        w=Inches(7.18),
        h=Inches(4.68),
        caption="从含水印图像中恢复 payload，再交给水印解码器重建 watermark Y",
        slide_title="关键模块：载荷恢复模块",
        placeholder_name="载荷恢复模块",
        needed_type="载荷恢复模块结构图",
        suggestion="若有高分辨率矢量版，优先替换 Word 导出的图片",
        missing_assets=missing_assets,
    )
    add_rect(slide, Inches(8.66), Inches(1.16), Inches(3.82), Inches(5.28))
    add_textbox(slide, Inches(8.88), Inches(1.40), Inches(3.36), Inches(0.30), "模块作用", font_size=16, bold=True, color=ACCENT)
    add_formula_card(slide, Inches(8.88), Inches(1.94), Inches(3.36), Inches(1.10), "恢复", "P_hat = f_rev(C_w)", "从 container 中估计 payload。", accent=ACCENT, formula_size=15.5)
    add_formula_card(slide, Inches(8.88), Inches(3.22), Inches(3.36), Inches(1.10), "攻击后", "P_hat_a = f_rev(A(C_w))", "攻击路径下仍尝试恢复主要 payload 结构。", accent=ACCENT_2, formula_size=12.8)
    add_card(slide, Inches(8.88), Inches(4.62), Inches(3.36), Inches(1.08), "设计重点", "恢复链路先保证 payload 稳定，再通过解码器恢复水印，便于分别约束中间表征和最终水印质量。", accent=ACCENT_3)
    set_notes(slide, "第三个关键模块是载荷恢复模块。它负责从含水印图像中估计 payload，然后再交给后面的水印解码器恢复 watermark Y。\n\n这个模块对鲁棒性非常关键。因为真实传播过程中，图像可能已经经过压缩、裁剪、缩放、模糊或者马赛克处理，如果恢复模块不能稳定取回 payload，后面再强的解码器也很难恢复出水印。因此本文在训练和损失设计中都对 payload 的恢复进行了单独约束。")
    slide_summaries.append(("关键模块：载荷恢复模块", ["单模块大图", "恢复公式", "攻击后恢复说明"]))

    # 10. Robust training flow
    slide = acquire_slide(prs, 10, target_slides, update_existing)
    set_bg(slide)
    add_title(slide, "鲁棒训练流程", "Robust training path")
    add_rect(slide, Inches(0.78), Inches(1.16), Inches(11.82), Inches(4.96), fill_color=CARD, line_color=LINE, radius=False)
    add_picture_or_placeholder(
        slide,
        path=media_path("image13.png"),
        x=Inches(0.96),
        y=Inches(1.36),
        w=Inches(11.46),
        h=Inches(4.52),
        caption="攻击仿真与鲁棒训练路径结构图",
        slide_title="鲁棒训练流程",
        placeholder_name="攻击仿真与鲁棒路径图",
        needed_type="裁剪/JPEG/缩放/模糊/马赛克攻击训练示意图",
        suggestion="若有原始 SVG，替换后放映效果会更好",
        missing_assets=missing_assets,
    )
    add_card(slide, Inches(0.82), Inches(6.22), Inches(2.18), Inches(0.78), "Crop", "随机裁剪用于模拟局部缺失。", accent=ACCENT)
    add_card(slide, Inches(3.14), Inches(6.22), Inches(2.18), Inches(0.78), "JPEG", "压缩失真是最常见传播扰动。", accent=ACCENT_2)
    add_card(slide, Inches(5.46), Inches(6.22), Inches(2.18), Inches(0.78), "Resize", "缩放会破坏细节与频率分布。", accent=ACCENT_3)
    add_card(slide, Inches(7.78), Inches(6.22), Inches(2.18), Inches(0.78), "Blur", "模糊会削弱恢复阶段的结构线索。", accent=ACCENT)
    add_card(slide, Inches(10.10), Inches(6.22), Inches(2.18), Inches(0.78), "Mosaic", "马赛克对应局部块状信息退化。", accent=ACCENT_2)
    set_notes(slide, "为了让模型更接近真实传播场景，训练阶段加入了多种攻击增强，包括裁剪、JPEG 压缩、缩放、模糊和马赛克。这样模型在训练时就能接触到常见失真，而不是只在 clean 图像上优化。\n\n这种鲁棒训练会带来一个权衡：clean 条件下的恢复质量可能会略有下降，但攻击后的稳定性会有所提升。后面的对比实验也会围绕这个权衡展开。")
    slide_summaries.append(("鲁棒训练流程", ["单页单流程图", "五类攻击增强说明"]))

    # 11. Loss design
    slide = acquire_slide(prs, 11, target_slides, update_existing)
    set_bg(slide)
    add_title(slide, "损失函数设计", "Loss decomposition")
    add_formula_card(
        slide,
        Inches(0.78),
        Inches(1.24),
        Inches(11.82),
        Inches(0.96),
        "总目标",
        "L_total = lambda1 (L_host + L_hvs) + lambda2 L_payload + lambda3 L_wm + lambda4 L_noise",
        "总目标同时约束不可见性、结构感知、payload 可恢复性和噪声分布稳定性。",
        accent=ACCENT,
        formula_size=15.0,
        note_size=10.0,
    )
    add_formula_card(
        slide,
        Inches(0.78),
        Inches(2.48),
        Inches(3.72),
        Inches(1.36),
        "Host Loss",
        "L_host = MSE(C, C_w)",
        "直接用 MSE 约束 cover 与 container 接近，是最基础的不可见性项。",
        accent=ACCENT,
        formula_size=14.8,
    )
    add_formula_card(
        slide,
        Inches(4.84),
        Inches(2.48),
        Inches(3.72),
        Inches(1.36),
        "HVS Loss",
        "L_hvs = mean(w * (C - C_w)^2)",
        "在平坦区域给予更高惩罚，让嵌入更符合人眼视觉敏感性。",
        accent=ACCENT_3,
        formula_size=13.6,
    )
    add_formula_card(
        slide,
        Inches(8.90),
        Inches(2.48),
        Inches(3.72),
        Inches(1.36),
        "Payload Loss",
        "L_payload = MSE(P, P_hat)",
        "保证恢复链路先能把中间 payload 稳定取回。",
        accent=ACCENT_2,
        formula_size=14.6,
    )
    add_formula_card(
        slide,
        Inches(0.78),
        Inches(4.14),
        Inches(5.70),
        Inches(1.58),
        "Watermark Loss",
        "L_wm = MSE(W_Y, W_hat_Y) + beta * (1 - SSIM(W_Y, W_hat_Y))",
        "同时约束像素误差和结构相似性，避免只得到模糊但数值接近的恢复结果。",
        accent=ACCENT,
        formula_size=12.2,
        note_size=9.2,
    )
    add_formula_card(
        slide,
        Inches(6.90),
        Inches(4.14),
        Inches(5.70),
        Inches(1.58),
        "Noise Regularization",
        "L_noise = mean(P)^2 + (std(P) - sigma0)^2 + Corr_adj(P)",
        "限制 payload 的均值、方差和局部相关性，使其更接近噪声化分布。",
        accent=ACCENT_2,
        formula_size=11.8,
        note_size=9.2,
    )
    add_rect(slide, Inches(0.78), Inches(6.00), Inches(11.82), Inches(0.64), fill_color=rgb("EEF5FF"), line_color=LINE)
    add_textbox(slide, Inches(0.98), Inches(6.18), Inches(11.4), Inches(0.22), "整体目标：在保持宿主图像视觉质量的同时，提高 payload 与 watermark Y 的稳定恢复能力。", font_size=10.8, color=MUTED)
    set_notes(slide, "损失函数的设计对应本文的三个目标：宿主图像要尽量不失真，payload 要能稳定恢复，最终水印结构也要尽量接近原始水印。\n\n总损失里，Host Loss 和 HVS Loss 主要约束不可见性，其中 HVS Loss 会让模型在视觉敏感区域更加谨慎；Payload Loss 约束中间载荷的一致性；Watermark Loss 同时使用 MSE 和 SSIM，兼顾像素误差和结构相似性；Noise Regularization 用来约束 payload 的统计分布，避免隐写信号过于集中或异常。")
    slide_summaries.append(("损失函数设计", ["总损失公式", "五类子项公式卡"]))

    # 12. Experiment setup
    slide = acquire_slide(prs, 12, target_slides, update_existing)
    set_bg(slide)
    add_title(slide, "实验设置与复现链路", "Implementation and evaluation")
    add_card(slide, Inches(0.75), Inches(1.2), Inches(2.7), Inches(1.15), "数据集", "训练集使用 DIV2K；测试集为 `inputs/chaos_batch` 中的自建图像样本。", accent=ACCENT)
    add_card(slide, Inches(3.62), Inches(1.2), Inches(2.7), Inches(1.15), "实现框架", "PyTorch + 自定义攻击模块，代码已包含训练、推理和批量测试流程。", accent=ACCENT_2)
    add_card(slide, Inches(6.49), Inches(1.2), Inches(2.7), Inches(1.15), "评价对象", "Host 指标用于衡量不可见性；WM 指标用于衡量水印恢复质量。", accent=ACCENT_3)
    add_card(slide, Inches(9.36), Inches(1.2), Inches(3.05), Inches(1.15), "实验类型", "主实验、Plain vs Robust 对比、A0~A5 消融，以及攻击后样例可视化。", accent=ACCENT)
    add_rect(slide, Inches(0.75), Inches(2.62), Inches(5.28), Inches(4.0))
    add_textbox(slide, Inches(0.95), Inches(2.82), Inches(4.8), Inches(0.3), "代码链路", font_size=15, bold=True, color=MUTED)
    add_bullets(
        slide,
        Inches(0.95),
        Inches(3.18),
        Inches(4.7),
        Inches(2.95),
        [
            "`experiments/edge_wm_noise_train.py`：训练主脚本",
            "`experiments/edge_wm_noise_infer.py`：单图推理与可视化导出",
            "`experiments/edge_wm_noise_test.py`：批量测试与指标汇总",
            "`scripts/run_compare_robust_vs_plain.sh`：对比实验",
            "`scripts/run_ablation_suite.sh`：消融实验",
            "`app_demo.py`：交互演示程序"
        ],
        font_size=13.9,
    )
    add_picture_or_placeholder(
        slide,
        path=fig_path("fig_main_model_training_readiness_zh.png"),
        x=Inches(6.28),
        y=Inches(2.62),
        w=Inches(6.24),
        h=Inches(4.0),
        caption="主实验模型训练到位性诊断图",
        slide_title="实验设置与复现链路",
        placeholder_name="训练到位性图",
        needed_type="训练过程诊断图",
        suggestion="当前图可直接使用，无需额外素材",
        missing_assets=missing_assets,
    )
    set_notes(slide, "实验部分使用 DIV2K 作为训练集，自建图像样本作为测试集。评价时分为两类对象：一类是宿主图像质量，也就是 container 相对于 cover 的失真程度；另一类是水印恢复质量，也就是恢复水印和原始 watermark Y 的接近程度。\n\n工程实现上，项目包含训练、单图推理、批量测试、对比实验、消融实验和 Streamlit 演示脚本。右侧训练到位性图用于说明主实验模型已经基本收敛，后续指标具有比较意义。")
    slide_summaries.append(("实验设置与复现链路", ["四项实验设置", "代码链路", "训练到位性图"]))

    # 13. Metrics intro
    slide = acquire_slide(prs, 13, target_slides, update_existing)
    set_bg(slide)
    add_title(slide, "评价指标说明", "MSE, PSNR and SSIM")
    add_rect(slide, Inches(0.78), Inches(1.16), Inches(11.82), Inches(0.72), fill_color=rgb("EEF5FF"), line_color=LINE)
    add_textbox(slide, Inches(0.98), Inches(1.33), Inches(11.4), Inches(0.26), "MSE、PSNR 和 SSIM 分别从像素误差、整体失真和结构相似性三个角度评价图像质量。", font_size=12.2, color=INK)
    add_formula_card(
        slide,
        Inches(0.78),
        Inches(2.06),
        Inches(3.72),
        Inches(2.92),
        "MSE",
        "MSE = (1 / (H * W)) * sum (I - I_hat)^2",
        "衡量像素级误差。数值越小，说明两张图在逐像素层面越接近；它也是 PSNR 的基础。",
        accent=ACCENT,
        formula_size=11.6,
        note_size=9.2,
    )
    add_formula_card(
        slide,
        Inches(4.80),
        Inches(2.06),
        Inches(3.72),
        Inches(2.92),
        "PSNR",
        "PSNR = 10 log10(MAX^2 / MSE)",
        "把误差转换成分贝尺度。数值越高，通常表示失真越小，更适合横向比较图像质量。",
        accent=ACCENT_2,
        formula_size=12.4,
        note_size=9.2,
    )
    add_formula_card(
        slide,
        Inches(8.82),
        Inches(2.06),
        Inches(3.72),
        Inches(2.92),
        "SSIM",
        "SSIM = l(I,I_hat) * c(I,I_hat) * s(I,I_hat)",
        "衡量亮度、对比度和结构一致性。它比单独看像素误差更接近人眼主观感受。",
        accent=ACCENT_3,
        formula_size=11.6,
        note_size=9.2,
    )
    add_card(slide, Inches(0.78), Inches(5.38), Inches(3.72), Inches(1.02), "Host 指标", "比较 cover 与 container，用来说明水印嵌入后宿主图像是否仍然自然、不可见。", accent=ACCENT)
    add_card(slide, Inches(4.80), Inches(5.38), Inches(3.72), Inches(1.02), "WM 指标", "比较 watermark Y 与 recovered watermark Y，用来说明水印亮度结构是否被稳定取回。", accent=ACCENT_2)
    add_card(slide, Inches(8.82), Inches(5.38), Inches(3.72), Inches(1.02), "Attack 指标", "在攻击后重新计算 WM PSNR/SSIM，用来体现方法对压缩、裁剪和模糊等扰动的鲁棒性。", accent=ACCENT_3)
    set_notes(slide, "这里简单说明三个评价指标。MSE 是均方误差，反映两张图在像素层面的平均差异，数值越小越好。PSNR 由 MSE 推出，单位是 dB，一般来说 PSNR 越高，说明图像失真越小。\n\nSSIM 更关注亮度、对比度和结构相似性，比单纯的像素误差更接近人眼感受。在本文中，Host PSNR 和 Host SSIM 用来评价含水印图像是否自然；WM PSNR 和 WM SSIM 用来评价水印是否恢复成功；攻击后的 WM 指标则用来评价鲁棒性。")
    slide_summaries.append(("评价指标说明", ["MSE/PSNR/SSIM 笼统说明", "Host/WM/Attack 指标用途"]))

    # 14. Qualitative results
    slide = acquire_slide(prs, 14, target_slides, update_existing)
    set_bg(slide)
    add_title(slide, "定性结果：嵌入、恢复与攻击", "Qualitative visualization")
    add_picture_or_placeholder(
        slide,
        path=media_path("image15.png"),
        x=Inches(0.78),
        y=Inches(1.28),
        w=Inches(5.85),
        h=Inches(4.55),
        caption="图 6-1 水印嵌入与恢复效果对比",
        slide_title="定性结果：嵌入、恢复与攻击",
        placeholder_name="嵌入与恢复效果图",
        needed_type="cover/container/residual/watermark recovery 对比图",
        suggestion="当前 Word 图已可用；若要更大图，可用 `outputs/edge_wm_noise_test/.../samples` 重新拼版",
        missing_assets=missing_assets,
    )
    add_picture_or_placeholder(
        slide,
        path=media_path("image16.png"),
        x=Inches(6.73),
        y=Inches(1.28),
        w=Inches(5.82),
        h=Inches(4.55),
        caption="图 6-2 常见攻击后的恢复效果",
        slide_title="定性结果：嵌入、恢复与攻击",
        placeholder_name="各类攻击下的恢复效果图",
        needed_type="crop/jpeg/resize/blur/mosaic/all attack 可视化",
        suggestion="当前 Word 图已可用；若想放大某一攻击，可单独拼高清对比",
        missing_assets=missing_assets,
    )
    add_card(slide, Inches(0.86), Inches(6.05), Inches(5.78), Inches(0.82), "观察 1", "clean 条件下宿主图几乎看不到明显水印痕迹，说明 residual 嵌入较克制。", accent=ACCENT)
    add_card(slide, Inches(6.81), Inches(6.05), Inches(5.72), Inches(0.82), "观察 2", "攻击会降低恢复清晰度，但多数场景下仍保留主要轮廓，与定量指标趋势一致。", accent=ACCENT_3)
    set_notes(slide, "接下来先看定性结果。左侧是 clean 条件下的嵌入和恢复效果，可以看到含水印图像和原图在视觉上差异很小，同时恢复出的水印仍然保留了主要结构。\n\n右侧是不同攻击后的恢复效果。攻击会降低水印的清晰度，这是符合预期的，但多数情况下仍能看到主要轮廓。这里需要强调的是，本文恢复的是水印的 Y 通道亮度结构，所以评价重点放在结构是否稳定保留。")
    slide_summaries.append(("定性结果：嵌入、恢复与攻击", ["clean可视化", "attack可视化"]))

    # 15. Main results
    slide = acquire_slide(prs, 15, target_slides, update_existing)
    set_bg(slide)
    add_title(slide, "主实验结果", "Main experiment")
    add_metric_card(slide, Inches(0.82), Inches(1.35), Inches(1.95), Inches(0.95), "Clean Host PSNR", f"{main_clean['host_psnr']:.3f} dB", "宿主图像失真较小")
    add_metric_card(slide, Inches(2.96), Inches(1.35), Inches(1.95), Inches(0.95), "Clean Host SSIM", f"{main_clean['host_ssim']:.4f}", "结构保持较好")
    add_metric_card(slide, Inches(0.82), Inches(2.48), Inches(1.95), Inches(0.95), "Clean WM PSNR", f"{main_clean['wm_psnr']:.3f} dB", "无攻击时恢复质量较高")
    add_metric_card(slide, Inches(2.96), Inches(2.48), Inches(1.95), Inches(0.95), "Clean WM SSIM", f"{main_clean['wm_ssim']:.4f}", "亮度结构恢复稳定")
    add_metric_card(slide, Inches(0.82), Inches(3.61), Inches(1.95), Inches(0.95), "Attack WM PSNR", f"{main_attack['wm_psnr_attack']:.3f} dB", "攻击后明显下降")
    add_metric_card(slide, Inches(2.96), Inches(3.61), Inches(1.95), Inches(0.95), "Attack WM SSIM", f"{main_attack['wm_ssim_attack']:.4f}", "但仍保留部分结构")
    add_card(slide, Inches(0.82), Inches(5.0), Inches(4.1), Inches(1.15), "结论", "主实验说明该方法在不可见性与可恢复性之间取得了平衡；综合攻击会显著压低恢复质量，但模型仍有一定鲁棒性。", accent=ACCENT)
    add_picture_or_placeholder(
        slide,
        path=fig_path("fig_main_model_final_quality_zh.png"),
        x=Inches(5.36),
        y=Inches(1.35),
        w=Inches(7.0),
        h=Inches(2.2),
        caption="主实验模型最终测试质量对比",
        slide_title="主实验结果",
        placeholder_name="主实验最终质量图",
        needed_type="clean 与 attack-all 的关键指标对比图",
        suggestion="当前中文重绘版可直接使用",
        missing_assets=missing_assets,
    )
    add_picture_or_placeholder(
        slide,
        path=fig_path("fig_clean_vs_attack_all_formal_runs_zh.png"),
        x=Inches(5.36),
        y=Inches(3.88),
        w=Inches(7.0),
        h=Inches(2.2),
        caption="正式实验 clean 与 attack-all 水印指标对比",
        slide_title="主实验结果",
        placeholder_name="clean vs attack-all 对比图",
        needed_type="正式实验整体对比图",
        suggestion="当前中文重绘版可直接使用",
        missing_assets=missing_assets,
    )
    set_notes(slide, f"主实验结果主要看两组指标。首先是 clean 条件下，Host PSNR 为 {main_clean['host_psnr']:.3f} dB，Host SSIM 为 {main_clean['host_ssim']:.4f}，说明嵌入后宿主图像的失真较小；同时，WM PSNR 为 {main_clean['wm_psnr']:.3f} dB，WM SSIM 为 {main_clean['wm_ssim']:.4f}，说明无攻击时水印恢复质量较高。\n\n第二组是 attack-all 条件下的恢复指标。可以看到，综合攻击后 WM PSNR 下降到 {main_attack['wm_psnr_attack']:.3f} dB，WM SSIM 下降到 {main_attack['wm_ssim_attack']:.4f}。这说明多种攻击叠加会明显压低恢复质量，但模型仍然保留了一定的结构恢复能力。")
    slide_summaries.append(("主实验结果", ["六个关键指标", "主实验最终质量图", "clean vs attack-all 图"]))

    # 16. Compare
    slide = acquire_slide(prs, 16, target_slides, update_existing)
    set_bg(slide)
    add_title(slide, "对比实验：Plain vs Robust", "Comparison study")
    add_picture_or_placeholder(
        slide,
        path=fig_path("fig_compare_plain_vs_robust_metrics_zh.png"),
        x=Inches(0.78),
        y=Inches(1.28),
        w=Inches(7.2),
        h=Inches(4.45),
        caption="普通训练与鲁棒训练的 clean/attack 核心指标对比",
        slide_title="对比实验：Plain vs Robust",
        placeholder_name="Plain vs Robust 指标图",
        needed_type="两组实验的 clean/attack 指标条形图",
        suggestion="当前中文重绘版可直接使用",
        missing_assets=missing_assets,
    )
    add_picture_or_placeholder(
        slide,
        path=fig_path("fig_compare_attack_retention_zh.png"),
        x=Inches(8.2),
        y=Inches(1.28),
        w=Inches(4.2),
        h=Inches(4.45),
        caption="攻击后水印质量保持率",
        slide_title="对比实验：Plain vs Robust",
        placeholder_name="攻击保持率图",
        needed_type="attack retention 图",
        suggestion="当前中文重绘版可直接使用",
        missing_assets=missing_assets,
    )
    add_card(
        slide,
        Inches(0.88),
        Inches(6.0),
        Inches(11.5),
        Inches(0.82),
        "结论",
        f"Plain 在 clean 条件下更优（Host PSNR {compare['plain']['clean']['host_psnr']:.3f} dB，WM PSNR {compare['plain']['clean']['wm_psnr']:.3f} dB）；Robust 在 attack WM PSNR 上略高（{compare['robust']['attack_all']['wm_psnr_attack']:.3f} dB vs {compare['plain']['attack_all']['wm_psnr_attack']:.3f} dB），体现了鲁棒训练的性能折中。",
        accent=ACCENT_3,
    )
    set_notes(slide, f"对比实验比较的是普通训练和鲁棒训练。结果可以概括为一个权衡：Plain 在 clean 条件下通常更好，例如 WM PSNR 为 {compare['plain']['clean']['wm_psnr']:.3f} dB，而 Robust 的 clean WM PSNR 为 {compare['robust']['clean']['wm_psnr']:.3f} dB。\n\n但是在 attack-all 条件下，Robust 的 attack WM PSNR 为 {compare['robust']['attack_all']['wm_psnr_attack']:.3f} dB，略高于 Plain 的 {compare['plain']['attack_all']['wm_psnr_attack']:.3f} dB。也就是说，鲁棒训练并不是让所有指标都提升，而是牺牲一部分 clean 表现，换取攻击场景下更稳定的恢复。")
    slide_summaries.append(("对比实验：Plain vs Robust", ["指标对比图", "attack retention 图", "性能折中结论"]))

    # 17. Ablation quality
    slide = acquire_slide(prs, 17, target_slides, update_existing)
    set_bg(slide)
    add_title(slide, "消融实验：恢复质量变化", "Ablation study")
    add_picture_or_placeholder(
        slide,
        path=fig_path("fig_ablation_wm_psnr_errorbar_zh.png"),
        x=Inches(0.78),
        y=Inches(1.28),
        w=Inches(5.85),
        h=Inches(4.5),
        caption="水印 PSNR 均值与 P10-P90 区间",
        slide_title="消融实验：恢复质量变化",
        placeholder_name="Ablation WM PSNR 图",
        needed_type="A0~A5 水印 PSNR 误差条图",
        suggestion="当前中文重绘版可直接使用",
        missing_assets=missing_assets,
    )
    add_picture_or_placeholder(
        slide,
        path=fig_path("fig_ablation_wm_ssim_errorbar_zh.png"),
        x=Inches(6.73),
        y=Inches(1.28),
        w=Inches(5.82),
        h=Inches(4.5),
        caption="水印 SSIM 均值与 P10-P90 区间",
        slide_title="消融实验：恢复质量变化",
        placeholder_name="Ablation WM SSIM 图",
        needed_type="A0~A5 水印 SSIM 误差条图",
        suggestion="当前中文重绘版可直接使用",
        missing_assets=missing_assets,
    )
    add_card(slide, Inches(0.86), Inches(5.98), Inches(3.8), Inches(0.84), "A0_Full", "整体最均衡，说明完整模块与损失组合是合理的。", accent=ACCENT)
    add_card(slide, Inches(4.76), Inches(5.98), Inches(3.8), Inches(0.84), "A1_NoRobust", "下降最明显，说明鲁棒训练不仅影响攻击场景，也影响整体恢复稳定性。", accent=ACCENT_3)
    add_card(slide, Inches(8.66), Inches(5.98), Inches(3.8), Inches(0.84), "A2_NoNoiseReg", "Attack-all 下 SSIM 下滑更明显，说明噪声正则有助于结构一致性。", accent=ACCENT_2)
    set_notes(slide, "消融实验用来说明各个设计是否有必要。这里比较的是 A0 到 A5 不同变体在水印恢复质量上的变化。\n\n从结果看，完整模型 A0_Full 的整体表现最均衡。去掉鲁棒训练的 A1_NoRobust 下降比较明显，说明鲁棒训练不仅影响攻击后的表现，也会影响整体恢复稳定性。A2_NoNoiseReg 在 attack-all 条件下的 SSIM 下降更明显，说明噪声正则对保持结构一致性也有帮助。")
    slide_summaries.append(("消融实验：恢复质量变化", ["Ablation PSNR 图", "Ablation SSIM 图", "A0/A1/A2 结论"]))

    # 18. Ablation attack + host
    slide = acquire_slide(prs, 18, target_slides, update_existing)
    set_bg(slide)
    add_title(slide, "消融实验：攻击收益与宿主质量", "Ablation follow-up")
    add_picture_or_placeholder(
        slide,
        path=fig_path("fig_ablation_delta_vs_a0_attack_zh.png"),
        x=Inches(0.78),
        y=Inches(1.28),
        w=Inches(6.0),
        h=Inches(4.45),
        caption="相对 A0_Full 的 attack 指标变化",
        slide_title="消融实验：攻击收益与宿主质量",
        placeholder_name="Ablation delta vs A0 图",
        needed_type="各消融项相对 A0 的攻击指标变化图",
        suggestion="当前中文重绘版可直接使用",
        missing_assets=missing_assets,
    )
    add_picture_or_placeholder(
        slide,
        path=fig_path("fig_ablation_host_quality_zh.png"),
        x=Inches(6.95),
        y=Inches(1.28),
        w=Inches(5.6),
        h=Inches(4.45),
        caption="各消融变体的宿主图像质量",
        slide_title="消融实验：攻击收益与宿主质量",
        placeholder_name="Ablation host quality 图",
        needed_type="Host PSNR/SSIM 对比图",
        suggestion="当前中文重绘版可直接使用",
        missing_assets=missing_assets,
    )
    add_bullets(
        slide,
        Inches(0.95),
        Inches(5.95),
        Inches(11.4),
        Inches(0.85),
        [
            "A3_NoHVS 在 clean 水印恢复上较高，但并没有带来更好的综合表现。",
            "A4_NoPayload 的 Host SSIM 很高，但 watermark 恢复明显失衡，说明 payload 约束不能去掉。",
            "A5_NoSSIMinWM 的结构相似性下降，说明水印恢复端仍需要 SSIM 监督。"
        ],
        font_size=14.5,
    )
    set_notes(slide, "这一组结果进一步说明几个损失项的作用。A3_NoHVS 在 clean 条件下水印恢复指标较高，但综合表现没有超过完整模型，说明只追求恢复指标并不一定能得到最均衡的结果。\n\nA4_NoPayload 的宿主图像质量看起来较好，但 watermark 恢复明显失衡，说明 payload 约束不能去掉。A5_NoSSIMinWM 的结构相似性下降，说明水印恢复端仍然需要 SSIM 监督。综合来看，A0_Full 是更均衡的配置。")
    slide_summaries.append(("消融实验：攻击收益与宿主质量", ["相对A0攻击变化图", "宿主质量图", "A3/A4/A5 结论"]))

    # 19. App demo
    slide = acquire_slide(prs, 19, target_slides, update_existing)
    set_bg(slide)
    add_title(slide, "App 演示方案", "Streamlit demo")
    add_video_or_placeholder(
        slide,
        video_path=demo_video,
        poster_path=demo_video_poster,
        x=Inches(0.78),
        y=Inches(1.28),
        w=Inches(4.35),
        h=Inches(3.15),
        caption="Streamlit 演示视频",
        slide_title="App 演示方案",
        placeholder_name="Streamlit 演示视频",
        needed_type="16:9 横向 mp4/mov/webm，建议 720p 以上、15~40 秒",
        suggestion="优先录制一次完整演示：上传图片 -> 选择攻击模式 -> 显示恢复结果；默认读取 `outputs/defense_ppt/streamlit_demo.mp4`",
        missing_assets=missing_assets,
    )
    add_rect(slide, Inches(0.78), Inches(4.92), Inches(4.35), Inches(1.50))
    add_textbox(slide, Inches(0.95), Inches(5.05), Inches(3.95), Inches(0.22), "运行命令", font_size=14, bold=True, color=MUTED)
    add_textbox(slide, Inches(0.95), Inches(5.3), Inches(3.95), Inches(0.35), "bash scripts/with_conda_gp.sh streamlit run app_demo.py", font_size=12.5, color=INK)
    add_bullets(
        slide,
        Inches(0.95),
        Inches(5.66),
        Inches(3.95),
        Inches(0.7),
        [
            "上传载体图片",
            "选择 checkpoint 和水印模板",
            "设置攻击模式并开始演示",
            "观察含水印图、攻击后图和恢复水印"
        ],
        font_size=12.2,
    )
    demo_specs = [
        ("原始图片", demo_path("cover.png"), Inches(5.46), Inches(1.28)),
        ("含水印图", demo_path("container.png"), Inches(8.34), Inches(1.28)),
        ("攻击后图", demo_path("container_attacked.png"), Inches(5.46), Inches(4.02)),
        ("攻击后恢复水印", demo_path("wm_hat_attacked.png"), Inches(8.34), Inches(4.02)),
    ]
    for caption, path, x, y in demo_specs:
        add_picture_or_placeholder(
            slide,
            path=path,
            x=x,
            y=y,
            w=Inches(2.55),
            h=Inches(2.2),
            caption=caption,
            slide_title="App 演示方案",
            placeholder_name=caption,
            needed_type=f"{caption} 示例图",
            suggestion="当前工程样例可直接使用",
            missing_assets=missing_assets,
        )
    set_notes(slide, "系统演示部分基于现有 checkpoint，不需要重新训练模型。运行左下角的命令后，可以在 Streamlit 页面中上传载体图片，选择水印模板和攻击模式，然后查看含水印图像、攻击后图像以及恢复水印结果。\n\n如果现场环境允许，我会直接播放或操作演示视频；如果现场设备不稳定，右侧的四张样例图也可以作为演示结果的备用展示。这个系统主要用于说明本文方法已经形成了从模型到应用展示的完整链路。")
    slide_summaries.append(("App 演示方案", ["演示视频位", "启动命令", "四张演示样例图"]))

    # 20. Conclusion
    slide = acquire_slide(prs, 20, target_slides, update_existing)
    set_bg(slide)
    add_title(slide, "结论与展望", "Conclusion")
    add_card(slide, Inches(0.78), Inches(1.28), Inches(3.0), Inches(1.25), "结论 1", "Y 通道噪声载荷 + RGB 残差嵌入是一条可行路线。", accent=ACCENT)
    add_card(slide, Inches(3.98), Inches(1.28), Inches(3.0), Inches(1.25), "结论 2", "主实验表明模型在不可见性与恢复质量之间取得了较好平衡。", accent=ACCENT_2)
    add_card(slide, Inches(7.18), Inches(1.28), Inches(3.0), Inches(1.25), "结论 3", "鲁棒训练和复合损失能提升复杂攻击下的稳定恢复能力。", accent=ACCENT_3)
    add_card(slide, Inches(10.38), Inches(1.28), Inches(2.2), Inches(1.25), "结论 4", "工程链路完整，可支撑复现与系统展示。", accent=ACCENT)
    add_rect(slide, Inches(0.78), Inches(2.85), Inches(5.72), Inches(3.55))
    add_textbox(slide, Inches(0.98), Inches(3.03), Inches(5.2), Inches(0.28), "当前不足", font_size=15, bold=True, color=MUTED)
    add_bullets(
        slide,
        Inches(0.98),
        Inches(3.35),
        Inches(5.1),
        Inches(2.6),
        [
            "强裁剪、强压缩和多攻击叠加下，恢复质量仍会明显下降。",
            "当前主要围绕固定模板和固定分辨率进行实验，泛化能力仍需扩展。",
            "指标仍以 PSNR/SSIM 为主，后续可加入误码率和主观评价。"
        ],
        font_size=15.5,
    )
    add_rect(slide, Inches(6.78), Inches(2.85), Inches(5.8), Inches(3.55))
    add_textbox(slide, Inches(6.98), Inches(3.03), Inches(5.2), Inches(0.28), "后续工作", font_size=15, bold=True, color=MUTED)
    add_bullets(
        slide,
        Inches(6.98),
        Inches(3.35),
        Inches(5.1),
        Inches(2.6),
        [
            "更真实的攻击仿真：屏摄、旋转、透视变换、链式压缩。",
            "更强的鲁棒训练策略：课程学习、自适应攻击采样、对抗式增强。",
            "更广的工程目标：多模板、多载荷、可变分辨率和真实应用部署。"
        ],
        font_size=15.5,
    )
    set_notes(slide, "最后对本文工作做一个总结。本文提出了一种基于 Y 通道噪声载荷和 RGB 残差嵌入的不可见水印方法，并通过主实验、对比实验和消融实验验证了方法的有效性。\n\n从结果来看，模型可以在保持较高宿主图像质量的同时，实现较稳定的水印恢复；鲁棒训练和复合损失也能提升复杂攻击下的恢复稳定性。当然，当前方法在强裁剪、强压缩和多攻击叠加下仍然会出现明显退化。后续可以从更真实的攻击仿真、更强的鲁棒训练策略，以及多模板、多载荷和可变分辨率部署等方向继续改进。")
    slide_summaries.append(("结论与展望", ["四个结论", "当前不足", "后续工作"]))

    # 21. Q&A
    slide = acquire_slide(prs, 21, target_slides, update_existing)
    set_bg(slide, rgb("F0F5FB"))
    add_textbox(slide, Inches(0.95), Inches(1.72), Inches(11.3), Inches(0.8), "答辩完毕", font_size=30, bold=True, color=INK, align=PP_ALIGN.CENTER, font_name=FONT_HEAD)
    add_textbox(slide, Inches(0.95), Inches(2.56), Inches(11.3), Inches(0.52), "感谢各位老师指导，欢迎批评指正", font_size=18, color=MUTED, align=PP_ALIGN.CENTER)
    add_rect(slide, Inches(3.2), Inches(3.42), Inches(6.9), Inches(1.45), fill_color=CARD, line_color=LINE)
    add_textbox(slide, Inches(3.45), Inches(3.73), Inches(6.4), Inches(0.8), "Q  &  A", font_size=28, bold=True, color=ACCENT, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    set_notes(slide, "以上就是我的毕业设计汇报。感谢各位老师的聆听，恳请各位老师批评指正。")
    slide_summaries.append(("答辩完毕", ["Q&A 结束页"]))

    return slide_summaries


def write_missing_assets(path: Path, missing_assets: list[MissingAsset]) -> None:
    lines = ["# 缺图清单", ""]
    if not missing_assets:
        lines.append("当前 PPT 未出现缺失图片占位。")
        path.write_text("\n".join(lines), encoding="utf-8")
        return

    dedup: list[MissingAsset] = []
    seen: set[tuple[str, str]] = set()
    for item in missing_assets:
        key = (item.slide_title, item.placeholder)
        if key in seen:
            continue
        seen.add(key)
        dedup.append(item)

    lines.append("| 页标题 | 占位项 | 需要的类型 | 建议来源 |")
    lines.append("|---|---|---|---|")
    for item in dedup:
        lines.append(f"| {item.slide_title} | {item.placeholder} | {item.needed_type} | {item.suggestion} |")
    path.write_text("\n".join(lines), encoding="utf-8")


def write_slide_summary(path: Path, slide_summaries: list[tuple[str, list[str]]]) -> None:
    lines = ["# Slides Summary", ""]
    for idx, (title, parts) in enumerate(slide_summaries, start=1):
        lines.append(f"## {idx}. {title}")
        for part in parts:
            lines.append(f"- {part}")
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate or partially update the defense PPT.")
    parser.add_argument(
        "--slides",
        type=str,
        default="",
        help="Comma-separated slide numbers to update in-place, e.g. 9,13. Leave empty to rebuild the full PPT.",
    )
    return parser.parse_args()


def parse_slide_set(value: str) -> set[int]:
    if not value.strip():
        return set()
    slides: set[int] = set()
    for part in value.split(","):
        item = part.strip()
        if not item:
            continue
        n = int(item)
        if n < 1 or n > TOTAL_SLIDES:
            raise ValueError(f"slide number out of range: {n}")
        slides.add(n)
    return slides


def main() -> None:
    args = parse_args()
    target_slides = parse_slide_set(args.slides)

    ensure_dirs()
    media = extract_docx_media(DOCX_PATH, DOCX_MEDIA_DIR)
    metrics = load_metrics()
    missing_assets: list[MissingAsset] = []

    if target_slides:
        if not PPTX_PATH.exists():
            raise FileNotFoundError(f"existing ppt not found for partial update: {PPTX_PATH}")
        prs = Presentation(str(PPTX_PATH))
        ensure_slide_capacity(prs, TOTAL_SLIDES)
        slide_summaries = build_presentation(
            prs,
            media,
            metrics,
            missing_assets,
            target_slides=target_slides,
            update_existing=True,
        )
    else:
        prs = make_prs()
        slide_summaries = build_presentation(prs, media, metrics, missing_assets)
    prs.save(str(PPTX_PATH))

    if target_slides:
        report_missing_assets: list[MissingAsset] = []
        report_prs = make_prs()
        report_slide_summaries = build_presentation(report_prs, media, metrics, report_missing_assets)
    else:
        report_missing_assets = missing_assets
        report_slide_summaries = slide_summaries

    write_missing_assets(MISSING_ASSETS_PATH, report_missing_assets)
    write_slide_summary(SLIDE_SUMMARY_PATH, report_slide_summaries)

    print(f"saved: {PPTX_PATH}")
    print(f"saved: {MISSING_ASSETS_PATH}")
    print(f"saved: {SLIDE_SUMMARY_PATH}")
    print(f"slides: {len(report_slide_summaries)}")


if __name__ == "__main__":
    main()
